import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ── Paleta ──────────────────────────────────────────────────────────────────
_SEC_HDR_BG = {
    "periodo":  "374151",
    "volumen":  "1D4ED8",
    "compras":  "0F766E",
    "doc_class": "0E7490",
    "framework": "BE123C",
    "tipo":     "065F46",
    "vigencia": "92400E",
    "cruce":    "7C2D12",
    "borrado":  "4C1D95",
}
_SEC_DATA_BG = {
    "periodo":  "F3F4F6",
    "volumen":  "DBEAFE",
    "compras":  "CCFBF1",
    "doc_class": "CFFAFE",
    "framework": "FFE4E6",
    "tipo":     "D1FAE5",
    "vigencia": "FEF3C7",
    "cruce":    "FFEDD5",
    "borrado":  "EDE9FE",
}
_DOC_CLASS_DESCRIPTIONS = {
    "FO": "Pedido Marco",
    "NB": "Pedido Estándar",
    "UD": "Stock transport order",
    "ZADI": "Ctto./ OC Asign. directa",
    "ZAUT": "Pedido Automático",
    "ZIMP": "OC para importación",
    "ZLVM": "Pedido bajo valor mats.",
    "ZLVS": "Pedido bajo valor servicios",
    "ZPD1": "Donaciones",
    "ZPDO": "Pedido donac. asig. Dir.",
    "ZSCM": "Contrato menor",
    "ZSCS": "Servicio Estándar",
    "ZSEM": "Ctto./OC Emergencia",
}
_SEC_LABELS = {
    "periodo":  "",
    "volumen":  "Volumen",
    "compras":  "Grupo de Compras y Monto",
    "doc_class": "Cl. documento de compras",
    "framework": "Contrato Marco",
    "tipo":     "Tipo de Posición  (tipo dominante por contrato)",
    "vigencia": "Vigencia del Contrato  (fecha de vencimiento más reciente por contrato)",
    "cruce":    "CRUCE  —  Tipo de Posición  ×  Año de Vencimiento",
    "borrado":  "Indicador de Borrado  (valor dominante por contrato)",
}
_SEP_BG    = "D1D5DB"
_TOTALS_BG = "1E3A5F"
_GRAY_CELL = "E9ECEF"   # celdas vacías en bloques adicionales

# Bloques adicionales (C, V, VACIO) que se replican bajo el bloque D
_EXTRA_BLOCKS = [
    ("C",  "c",    "TIPO C  —  Consignación"),
    ("V",  "v",    "TIPO V  —  Valor límite"),
    ("",   "vacio","SIN TIPO  —  Posición sin tipo asignado"),
]


def _fill(hex_bg: str) -> PatternFill:
    return PatternFill(start_color=hex_bg, end_color=hex_bg, fill_type="solid")


def _section_labels(
    date_section_label: str | None = None,
    cruce_section_label: str | None = None,
) -> dict[str, str]:
    labels = dict(_SEC_LABELS)
    if date_section_label is not None:
        labels["vigencia"] = date_section_label
    if cruce_section_label is not None:
        labels["cruce"] = cruce_section_label
    return labels


def _extract_period(source_file: str) -> str:
    stem = Path(source_file).stem
    m = re.search(r"\d{4}[-_]\d{4}", stem)
    return m.group().replace("-", "_") if m else stem


def _is_sep(sec: str) -> bool:
    return sec.startswith("_sep")


def _text_series(df: pd.DataFrame, column: str, default: str = "") -> pd.Series:
    if column not in df.columns:
        return pd.Series([default] * len(df), index=df.index, dtype="object")
    series = df[column].fillna(default).astype(str).str.strip()
    return series.mask(series.str.lower().isin({"nan", "none", "nat"}), default)


def _value_specs(values: list[str], prefix: str, empty_label: str) -> list[tuple[str, str, str]]:
    normalized = sorted({str(v).strip() for v in values}, key=lambda v: (v == "", v))
    return [
        (value, f"{prefix}_{idx:02d}", empty_label if value == "" else value)
        for idx, value in enumerate(normalized, 1)
    ]


def _contracts_table(subset: pd.DataFrame, date_column: str = "validity_end") -> pd.DataFrame:
    tmp = pd.DataFrame({
        "purchase_document": subset["purchase_document"].values,
        "date_value": pd.to_datetime(
            subset[date_column].values if date_column in subset.columns
            else [pd.NaT] * len(subset),
            errors="coerce",
        ),
        "position_type": (
            subset["position_type"].fillna("").astype(str).str.strip().values
            if "position_type" in subset.columns else [""] * len(subset)
        ),
        "deletion_flag": (
            subset["deletion_flag"].fillna("").astype(str).str.strip().values
            if "deletion_flag" in subset.columns else [""] * len(subset)
        ),
        "purchase_doc_class": _text_series(subset, "purchase_doc_class").values,
        "framework_contract": _text_series(subset, "framework_contract").values,
        "purchase_group": (
            subset["purchase_group"].fillna("SIN_GRUPO").astype(str).str.strip().values
            if "purchase_group" in subset.columns else ["SIN_GRUPO"] * len(subset)
        ),
        "estimated_value": pd.to_numeric(
            subset["estimated_value"].values if "estimated_value" in subset.columns
            else [0.0] * len(subset),
            errors="coerce",
        ),
    })
    tmp["purchase_group"] = tmp["purchase_group"].where(tmp["purchase_group"] != "", other="SIN_GRUPO")
    contracts = tmp.groupby("purchase_document").agg(
        max_date=("date_value", "max"),
        dom_type=("position_type", lambda s: s.mode().iloc[0] if not s.empty else ""),
        dom_flag=("deletion_flag", lambda s: s.mode().iloc[0] if not s.empty else ""),
        dom_doc_class=("purchase_doc_class", lambda s: s.mode().iloc[0] if not s.empty else ""),
        dom_framework_contract=("framework_contract", lambda s: s.mode().iloc[0] if not s.empty else ""),
        dom_purchase_group=("purchase_group", lambda s: s.mode().iloc[0] if not s.empty else "SIN_GRUPO"),
        estimated_value_sum=("estimated_value", "sum"),
    )
    contracts["yr"] = contracts["max_date"].dt.year
    return contracts


def _framework_stats(subset: pd.DataFrame, framework_contract_col: str) -> dict:
    purchase_document = _text_series(subset, "purchase_document")
    framework_contract = _text_series(subset, framework_contract_col)
    tmp = pd.DataFrame({
        "purchase_document": purchase_document,
        "framework_contract": framework_contract,
    })
    tmp = tmp[tmp["purchase_document"] != ""]
    total_po = int(tmp["purchase_document"].nunique())
    linked = tmp[tmp["framework_contract"] != ""]
    po_with_framework = int(linked["purchase_document"].nunique())
    framework_count = int(linked["framework_contract"].nunique())
    pct = round(po_with_framework / total_po * 100, 2) if total_po else 0.0
    return {
        "framework_contracts": framework_count,
        "po_with_framework": po_with_framework,
        "po_without_framework": max(total_po - po_with_framework, 0),
        "pct_po_with_framework": pct,
    }


def _period_stats(
    subset: pd.DataFrame,
    future_years: list[int],
    date_column: str = "validity_end",
    include_purchase_amount_section: bool = True,
    doc_class_specs: list[tuple[str, str, str]] | None = None,
    include_framework_section: bool = False,
    framework_contract_col: str = "framework_contract",
) -> dict:
    doc_class_specs = doc_class_specs or []
    c = _contracts_table(subset, date_column=date_column)
    stats: dict = {
        "nro_lineas":    len(subset),
        "nro_contratos": len(c),
        "pos_D":         int((c["dom_type"] == "D").sum()),
        "pos_C":         int((c["dom_type"] == "C").sum()),
        "pos_V":         int((c["dom_type"] == "V").sum()),
        "pos_VACIO":     int((c["dom_type"] == "").sum()),
        "yr_leq2025":    int((c["yr"] <= 2025).sum()),
    }
    if include_purchase_amount_section:
        stats["grupos_compra"] = int(c["dom_purchase_group"].nunique())
        stats["monto_total_usd"] = round(float(c["estimated_value_sum"].sum()), 2)

    for value, key, _ in doc_class_specs:
        stats[key] = int((c["dom_doc_class"] == value).sum())

    if include_framework_section:
        stats.update(_framework_stats(subset, framework_contract_col))

    for yr in future_years:
        stats[f"yr_{yr}"] = int((c["yr"] == yr).sum())
    stats["yr_sin_fecha"] = int(c["yr"].isna().sum())

    # Bloque D
    stats["d_leq2025"] = int(((c["dom_type"] == "D") & (c["yr"] <= 2025)).sum())
    for yr in future_years:
        stats[f"d_{yr}"] = int(((c["dom_type"] == "D") & (c["yr"] == yr)).sum())

    # Bloques C, V, VACIO
    for typ, pfx, _ in _EXTRA_BLOCKS:
        stats[f"{pfx}_leq2025"] = int(((c["dom_type"] == typ) & (c["yr"] <= 2025)).sum())
        for yr in future_years:
            stats[f"{pfx}_{yr}"] = int(((c["dom_type"] == typ) & (c["yr"] == yr)).sum())

    stats["flag_L"]   = int((c["dom_flag"] == "L").sum())
    stats["flag_S"]   = int((c["dom_flag"] == "S").sum())
    stats["flag_sin"] = int((c["dom_flag"] == "").sum())
    return stats


def _date_bucket_stats(
    subset: pd.DataFrame,
    future_years: list[int],
    date_column: str,
) -> dict:
    c = _contracts_table(subset, date_column=date_column)
    stats: dict = {
        "yr_leq2025": int((c["yr"] <= 2025).sum()),
    }
    for yr in future_years:
        stats[f"yr_{yr}"] = int((c["yr"] == yr).sum())
    stats["yr_sin_fecha"] = int(c["yr"].isna().sum())
    return stats


def _date_bucket_rows(
    df: pd.DataFrame,
    future_years: list[int],
    date_column: str,
    operation: str,
) -> list[dict]:
    period_rows: list[dict] = []
    if "_source_file" in df.columns:
        src_files = sorted(df["_source_file"].unique(), key=lambda s: _extract_period(s))
        for src_file in src_files:
            subset = df[df["_source_file"] == src_file]
            stats = _date_bucket_stats(subset, future_years, date_column)
            stats["periodo"] = _extract_period(str(src_file))
            period_rows.append(stats)
    else:
        stats = _date_bucket_stats(df, future_years, date_column)
        stats["periodo"] = operation
        period_rows.append(stats)

    total = _date_bucket_stats(df, future_years, date_column)
    total["periodo"] = "TOTAL"
    return period_rows + [total]


def _remap_cruce(row_data: dict, pfx: str, future_years: list[int]) -> dict:
    """Return row_data with d_* cruce keys replaced by pfx_* values."""
    remapped = dict(row_data)
    remapped["d_leq2025"] = row_data.get(f"{pfx}_leq2025", 0)
    for yr in future_years:
        remapped[f"d_{yr}"] = row_data.get(f"{pfx}_{yr}", 0)
    return remapped


def _col_defs(
    future_years: list[int],
    include_purchase_amount_section: bool = True,
    doc_class_specs: list[tuple[str, str, str]] | None = None,
    include_framework_section: bool = False,
    date_past_label: str = "Vencidos hasta 2025",
    date_future_label_prefix: str = "Vigentes en",
    date_empty_label: str = "Sin fecha de vencimiento",
    date_empty_note: str = "(validity_end vacío)",
) -> list[tuple[str, str, str, str]]:
    doc_class_specs = doc_class_specs or []
    d: list[tuple[str, str, str, str]] = []

    d += [("periodo", "periodo", "Período", "Archivo fuente")]
    d += [("_sep_0", "_sep_0", "", "")]

    d += [
        ("volumen", "nro_lineas",    "Total de posiciones",  "(filas en el dataset)"),
        ("volumen", "nro_contratos", "Contratos únicos",     "(Documentos de compra distintos)"),
    ]
    d += [("_sep_1", "_sep_1", "", "")]

    if include_purchase_amount_section:
        d += [
            ("compras", "grupos_compra", "Grupos de compras", "(grupos distintos)"),
            ("compras", "monto_total_usd", "Monto total USD", "(estimated_value)"),
        ]
        d += [("_sep_1b", "_sep_1b", "", "")]

    if doc_class_specs:
        for _, key, label in doc_class_specs:
            description = _DOC_CLASS_DESCRIPTIONS.get(label, "")
            d.append(("doc_class", key, f"Clase {label}", f"({description})" if description else ""))
        d += [("_sep_1c", "_sep_1c", "", "")]

    if include_framework_section:
        d += [
            ("framework", "framework_contracts", "Contratos Marco", "(distintos)"),
            ("framework", "po_with_framework", "PO con marco", "(documentos únicos)"),
            ("framework", "po_without_framework", "PO sin marco", "(documentos únicos)"),
            ("framework", "pct_po_with_framework", "% PO con marco", "(sobre documentos únicos)"),
        ]
        d += [("_sep_1d", "_sep_1d", "", "")]

    d += [
        ("tipo", "pos_D",     "Tipo D",             "Servicio / límite de valor"),
        ("tipo", "pos_C",     "Tipo C",             "Consignación"),
        ("tipo", "pos_V",     "Tipo V",             "Valor límite"),
        ("tipo", "pos_VACIO", "Sin tipo asignado",  "(position_type vacío)"),
    ]
    d += [("_sep_2", "_sep_2", "", "")]

    d += [("vigencia", "yr_leq2025", date_past_label,
           "(cualquier año ≤ 2025 — agrupados)")]
    for yr in future_years:
        d.append(("vigencia", f"yr_{yr}", f"{date_future_label_prefix} {yr}", f"(año {yr})"))
    d.append(("vigencia", "yr_sin_fecha", date_empty_label, date_empty_note))
    d += [("_sep_3", "_sep_3", "", "")]

    # Cruce — columnas genéricas (el bloque activo lo indica la fila de etiqueta)
    d.append(("cruce", "d_leq2025", f"× {date_past_label.replace('hasta', '≤')}",
              "(tipo según bloque — año agrupado)"))
    for yr in future_years:
        d.append(("cruce", f"d_{yr}", f"× {date_future_label_prefix} {yr}", f"(tipo según bloque — año {yr})"))
    d += [("_sep_4", "_sep_4", "", "")]

    d += [
        ("borrado", "flag_L",   "Libre — sin borrado (L)",   "(indicador L — no marcado)"),
        ("borrado", "flag_S",   "Marcado para borrar (S)",   "(indicador S — en proceso de baja)"),
        ("borrado", "flag_sin", "Sin indicador de borrado",  "(campo deletion_flag vacío)"),
    ]
    return d


def _write_section_headers(ws, col_defs: list, section_labels: dict[str, str] | None = None) -> None:
    section_labels = section_labels or _SEC_LABELS
    n = len(col_defs)
    groups: list[tuple[str, int, int]] = []
    cur_sec, start = col_defs[0][0], 1
    for i, (sec, *_) in enumerate(col_defs[1:], 2):
        if sec != cur_sec:
            groups.append((cur_sec, start, i - 1))
            cur_sec, start = sec, i
    groups.append((cur_sec, start, n))

    for sec, c_start, c_end in groups:
        if c_end > c_start:
            ws.merge_cells(start_row=2, start_column=c_start,
                           end_row=2, end_column=c_end)
        cell = ws.cell(row=2, column=c_start)
        if _is_sep(sec):
            cell.fill = _fill(_SEP_BG)
        else:
            cell.value = section_labels.get(sec, sec)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = _fill(_SEC_HDR_BG.get(sec, "374151"))
            cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_data_rows(
    ws,
    all_rows: list[dict],
    col_defs: list,
    start_row: int,
    only_cruce: bool = False,
) -> int:
    """Write data rows and return the next available row number."""
    row_n = start_row
    for row_data in all_rows:
        is_total = row_data["periodo"] == "TOTAL"
        for col_idx, (sec, key, *_) in enumerate(col_defs, 1):
            cell = ws.cell(row=row_n, column=col_idx)
            if _is_sep(sec):
                cell.fill = _fill(_SEP_BG)
                continue
            if key == "periodo":
                cell.value = row_data.get("periodo", "")
                if is_total:
                    cell.font = Font(bold=True, size=10, color="FFFFFF")
                    cell.fill = _fill(_TOTALS_BG)
                else:
                    cell.fill = _fill(_SEC_DATA_BG["periodo"])
                    cell.font = Font(bold=True, size=10)
                cell.alignment = Alignment(horizontal="left", vertical="center")
                continue
            if only_cruce and sec != "cruce":
                # En bloques adicionales, vaciar secciones que no son cruce
                cell.fill = _fill(_GRAY_CELL)
                continue
            val = row_data.get(key, 0)
            cell.value = val
            if is_total:
                cell.font = Font(bold=True, size=10, color="FFFFFF")
                cell.fill = _fill(_TOTALS_BG)
            else:
                cell.fill = _fill(_SEC_DATA_BG.get(sec, "FFFFFF"))
                cell.font = Font(bold=(key == "periodo"), size=10)
            cell.alignment = Alignment(
                horizontal="left" if key == "periodo" else "right",
                vertical="center",
            )
            if isinstance(val, int) and key != "periodo":
                cell.number_format = "#,##0"
            elif isinstance(val, float):
                cell.number_format = "#,##0.00"
        ws.row_dimensions[row_n].height = 18
        row_n += 1
    return row_n


def _write_block_label(ws, row_n: int, col_defs: list, label: str) -> int:
    """Label row: merges only the cruce-section columns; other cells get a neutral fill."""
    cruce_start = cruce_end = None
    for col_idx, (sec, *_) in enumerate(col_defs, 1):
        if sec == "cruce":
            if cruce_start is None:
                cruce_start = col_idx
            cruce_end = col_idx

    for col_idx, (sec, *_) in enumerate(col_defs, 1):
        cell = ws.cell(row=row_n, column=col_idx)
        if _is_sep(sec):
            cell.fill = _fill(_SEP_BG)
        elif sec == "cruce":
            cell.fill = _fill(_SEC_HDR_BG["cruce"])  # covered by merge below
        else:
            cell.fill = _fill(_GRAY_CELL)

    if cruce_start is not None and cruce_end is not None and cruce_end >= cruce_start:
        ws.merge_cells(start_row=row_n, start_column=cruce_start,
                       end_row=row_n, end_column=cruce_end)
    lc = ws.cell(row=row_n, column=cruce_start or 1,
                 value=f"  CRUCE  —  {label}")
    lc.font = Font(bold=True, size=10, color="FFFFFF")
    lc.fill = _fill(_SEC_HDR_BG["cruce"])
    lc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_n].height = 18
    return row_n + 1


def _write_section_block_label(ws, row_n: int, col_defs: list, section: str, label: str) -> int:
    section_start = section_end = None
    for col_idx, (sec, *_) in enumerate(col_defs, 1):
        if sec == section:
            if section_start is None:
                section_start = col_idx
            section_end = col_idx

    for col_idx, (sec, *_) in enumerate(col_defs, 1):
        cell = ws.cell(row=row_n, column=col_idx)
        if _is_sep(sec):
            cell.fill = _fill(_SEP_BG)
        elif sec == section:
            cell.fill = _fill(_SEC_HDR_BG.get(section, "374151"))
        else:
            cell.fill = _fill(_GRAY_CELL)

    if section_start is not None and section_end is not None and section_end >= section_start:
        ws.merge_cells(start_row=row_n, start_column=section_start,
                       end_row=row_n, end_column=section_end)
    cell = ws.cell(row=row_n, column=section_start or 1, value=f"  {label}")
    cell.font = Font(bold=True, size=10, color="FFFFFF")
    cell.fill = _fill(_SEC_HDR_BG.get(section, "374151"))
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_n].height = 18
    return row_n + 1


def _write_date_only_rows(ws, rows: list[dict], col_defs: list, start_row: int) -> int:
    row_n = start_row
    for row_data in rows:
        is_total = row_data["periodo"] == "TOTAL"
        for col_idx, (sec, key, *_) in enumerate(col_defs, 1):
            cell = ws.cell(row=row_n, column=col_idx)
            if _is_sep(sec):
                cell.fill = _fill(_SEP_BG)
                continue
            if key == "periodo":
                cell.value = row_data.get("periodo", "")
                if is_total:
                    cell.font = Font(bold=True, size=10, color="FFFFFF")
                    cell.fill = _fill(_TOTALS_BG)
                else:
                    cell.fill = _fill(_SEC_DATA_BG["periodo"])
                    cell.font = Font(bold=True, size=10)
                cell.alignment = Alignment(horizontal="left", vertical="center")
                continue
            if sec != "vigencia":
                cell.fill = _fill(_GRAY_CELL)
                continue
            cell.value = row_data.get(key, 0)
            if is_total:
                cell.font = Font(bold=True, size=10, color="FFFFFF")
                cell.fill = _fill(_TOTALS_BG)
            else:
                cell.font = Font(size=10)
                cell.fill = _fill(_SEC_DATA_BG["vigencia"])
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0"
        ws.row_dimensions[row_n].height = 18
        row_n += 1
    return row_n


def _overlay_section_block_label(ws, row_n: int, col_defs: list, section: str, label: str) -> None:
    section_start = section_end = None
    for col_idx, (sec, *_) in enumerate(col_defs, 1):
        if sec == section:
            if section_start is None:
                section_start = col_idx
            section_end = col_idx

    if section_start is None or section_end is None:
        return

    ws.merge_cells(start_row=row_n, start_column=section_start,
                   end_row=row_n, end_column=section_end)
    cell = ws.cell(row=row_n, column=section_start, value=f"  {label}")
    cell.font = Font(bold=True, size=10, color="FFFFFF")
    cell.fill = _fill(_SEC_HDR_BG.get(section, "374151"))
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _overlay_date_only_rows(ws, rows: list[dict], col_defs: list, start_row: int) -> None:
    for row_offset, row_data in enumerate(rows):
        row_n = start_row + row_offset
        is_total = row_data["periodo"] == "TOTAL"
        for col_idx, (sec, key, *_) in enumerate(col_defs, 1):
            if key != "periodo" and sec != "vigencia":
                continue
            cell = ws.cell(row=row_n, column=col_idx)
            if key == "periodo":
                cell.value = row_data.get("periodo", "")
                if is_total:
                    cell.font = Font(bold=True, size=10, color="FFFFFF")
                    cell.fill = _fill(_TOTALS_BG)
                else:
                    cell.fill = _fill(_SEC_DATA_BG["periodo"])
                    cell.font = Font(bold=True, size=10)
                cell.alignment = Alignment(horizontal="left", vertical="center")
                continue
            cell.value = row_data.get(key, 0)
            if is_total:
                cell.font = Font(bold=True, size=10, color="FFFFFF")
                cell.fill = _fill(_TOTALS_BG)
            else:
                cell.font = Font(size=10)
                cell.fill = _fill(_SEC_DATA_BG["vigencia"])
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0"


def _write_framework_detail_sheet(
    wb: Workbook,
    df: pd.DataFrame,
    framework_contract_col: str = "framework_contract",
) -> None:
    ws = wb.create_sheet("Contrato Marco")
    headers = [
        "Contrato Marco",
        "PO asociadas",
        "Posiciones",
        "Períodos fuente",
        "Clases doc. compra",
        "Primera fecha doc.",
        "Última fecha doc.",
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill(_SEC_HDR_BG["framework"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    if framework_contract_col not in df.columns or "purchase_document" not in df.columns:
        ws.cell(row=2, column=1, value="No existe una columna de Contrato Marco disponible para este dataset.")
        return

    tmp = pd.DataFrame({
        "framework_contract": _text_series(df, framework_contract_col),
        "purchase_document": _text_series(df, "purchase_document"),
        "position": _text_series(df, "position"),
        "source_file": _text_series(df, "_source_file"),
        "purchase_doc_class": _text_series(df, "purchase_doc_class"),
    })
    if "document_date" in df.columns:
        tmp["document_date"] = pd.to_datetime(df["document_date"], errors="coerce")
    else:
        tmp["document_date"] = pd.NaT

    tmp = tmp[(tmp["framework_contract"] != "") & (tmp["purchase_document"] != "")]
    if tmp.empty:
        ws.cell(row=2, column=1, value="No hay PO vinculadas a Contrato Marco en este dataset.")
        return

    rows: list[dict] = []
    for framework_contract, group in tmp.groupby("framework_contract", sort=True):
        doc_pos = group[["purchase_document", "position"]].drop_duplicates()
        periods = sorted({_extract_period(v) for v in group["source_file"] if v})
        doc_classes = sorted({v for v in group["purchase_doc_class"] if v})
        first_date = group["document_date"].min()
        last_date = group["document_date"].max()
        rows.append({
            "Contrato Marco": framework_contract,
            "PO asociadas": int(group["purchase_document"].nunique()),
            "Posiciones": int(len(doc_pos)),
            "Períodos fuente": ", ".join(periods),
            "Clases doc. compra": ", ".join(doc_classes),
            "Primera fecha doc.": "" if pd.isna(first_date) else first_date.strftime("%Y-%m-%d"),
            "Última fecha doc.": "" if pd.isna(last_date) else last_date.strftime("%Y-%m-%d"),
        })

    rows.sort(key=lambda r: (-r["PO asociadas"], str(r["Contrato Marco"])))
    for row_idx, row in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[header])
            cell.fill = _fill(_SEC_DATA_BG["framework"])
            cell.alignment = Alignment(
                horizontal="right" if isinstance(row[header], int) else "left",
                vertical="center",
            )
            if isinstance(row[header], int):
                cell.number_format = "#,##0"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    widths = [18, 14, 12, 24, 20, 16, 16]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def generate_stats_report(
    df: pd.DataFrame,
    operation: str,
    output_dir: Path,
    dataset_label: str = "Universo Completo — C1 (pre-reglas)",
    entity_label: str = "Contratos",
    file_suffix: str | None = None,
    date_column: str = "validity_end",
    date_section_label: str | None = None,
    cruce_section_label: str | None = None,
    date_past_label: str = "Vencidos hasta 2025",
    date_future_label_prefix: str = "Vigentes en",
    date_empty_label: str = "Sin fecha de vencimiento",
    date_empty_note: str = "(validity_end vacío)",
    include_delivery_date_section: bool = False,
    delivery_date_column: str = "delivery_date",
    delivery_date_label: str = "FECHA DE ENTREGA — fecha de entrega más reciente por documento",
    include_purchase_amount_section: bool = True,
    include_doc_class_section: bool = False,
    include_framework_section: bool = False,
    framework_contract_col: str = "framework_contract",
) -> Path:
    if date_column in df.columns:
        ve = pd.to_datetime(df[date_column], errors="coerce")
        future_years = sorted({int(y) for y in ve.dropna().dt.year.unique() if y >= 2026})
    else:
        future_years = []
    if include_delivery_date_section and delivery_date_column in df.columns:
        delivery_dates = pd.to_datetime(df[delivery_date_column], errors="coerce")
        delivery_years = {int(y) for y in delivery_dates.dropna().dt.year.unique() if y >= 2026}
        future_years = sorted(set(future_years) | delivery_years)

    doc_class_specs: list[tuple[str, str, str]] = []
    if include_doc_class_section:
        contracts = _contracts_table(df, date_column=date_column)
        doc_class_specs = _value_specs(
            contracts["dom_doc_class"].fillna("").astype(str).str.strip().tolist(),
            "doc_class",
            "Sin clase",
        )

    # Per-period stats — ordenados por período extraído (no por nombre de archivo)
    period_rows: list[dict] = []
    if "_source_file" in df.columns:
        src_files = sorted(df["_source_file"].unique(),
                           key=lambda s: _extract_period(s))
        for src_file in src_files:
            subset = df[df["_source_file"] == src_file]
            stats = _period_stats(
                subset,
                future_years,
                date_column=date_column,
                include_purchase_amount_section=include_purchase_amount_section,
                doc_class_specs=doc_class_specs,
                include_framework_section=include_framework_section,
                framework_contract_col=framework_contract_col,
            )
            stats["periodo"] = _extract_period(str(src_file))
            period_rows.append(stats)
    else:
        stats = _period_stats(
            df,
            future_years,
            date_column=date_column,
            include_purchase_amount_section=include_purchase_amount_section,
            doc_class_specs=doc_class_specs,
            include_framework_section=include_framework_section,
            framework_contract_col=framework_contract_col,
        )
        stats["periodo"] = operation
        period_rows.append(stats)

    numeric_keys = [k for k in period_rows[0] if k != "periodo"]
    total = {"periodo": "TOTAL"} | {k: sum(r[k] for r in period_rows) for k in numeric_keys}
    if include_purchase_amount_section and "purchase_group" in df.columns:
        purchase_group = df["purchase_group"].fillna("SIN_GRUPO").astype(str).str.strip()
        purchase_group = purchase_group.where(purchase_group != "", other="SIN_GRUPO")
        total["grupos_compra"] = int(purchase_group.nunique())
    if include_doc_class_section:
        contracts = _contracts_table(df, date_column=date_column)
        for value, key, _ in doc_class_specs:
            total[key] = int((contracts["dom_doc_class"] == value).sum())
    if include_framework_section:
        total.update(_framework_stats(df, framework_contract_col))
    all_rows = period_rows + [total]
    delivery_rows = (
        _date_bucket_rows(df, future_years, delivery_date_column, operation)
        if include_delivery_date_section and delivery_date_column in df.columns
        else []
    )

    col_defs = _col_defs(
        future_years,
        include_purchase_amount_section=include_purchase_amount_section,
        doc_class_specs=doc_class_specs,
        include_framework_section=include_framework_section,
        date_past_label=date_past_label,
        date_future_label_prefix=date_future_label_prefix,
        date_empty_label=date_empty_label,
        date_empty_note=date_empty_note,
    )
    ncols = len(col_defs)

    # ── Workbook ─────────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = f"Estadísticas {operation}"

    # Fila 1 — Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(row=1, column=1,
              value=(f"Estadísticas de {entity_label} — Operación {operation}   |   "
                       f"{dataset_label}   |   "
                       f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"))
    t.font = Font(bold=True, size=12, color="FFFFFF")
    t.fill = _fill("1E3A5F")
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    # Fila 2 — Secciones fusionadas
    _write_section_headers(
        ws,
        col_defs,
        _section_labels(
            date_section_label=date_section_label,
            cruce_section_label=cruce_section_label,
        ),
    )
    ws.row_dimensions[2].height = 20

    # Fila 3 — Sub-encabezados
    for col_idx, (sec, key, hdr1, hdr2) in enumerate(col_defs, 1):
        cell = ws.cell(row=3, column=col_idx)
        if _is_sep(sec):
            cell.fill = _fill(_SEP_BG)
        else:
            cell.value = f"{hdr1}\n{hdr2}"
            cell.font = Font(bold=True, size=9, color="FFFFFF")
            cell.fill = _fill(_SEC_HDR_BG.get(sec, "374151"))
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 42

    # ── Bloque D (completo, todas las columnas) ──────────────────────────────
    row_n = 4
    row_n = _write_block_label(ws, row_n, col_defs, "TIPO D  —  Servicio / límite de valor")
    row_n = _write_data_rows(ws, all_rows, col_defs, row_n, only_cruce=False)

    # ── Bloques C, V, VACIO (sólo cruce — otras columnas en gris) ────────────
    delivery_overlay_pending = bool(delivery_rows)
    for typ, pfx, label in _EXTRA_BLOCKS:
        # Fila en blanco
        for col_idx in range(1, ncols + 1):
            ws.cell(row=row_n, column=col_idx).fill = _fill("FFFFFF")
        ws.row_dimensions[row_n].height = 8
        row_n += 1

        # Etiqueta del bloque
        label_row = row_n
        row_n = _write_block_label(ws, row_n, col_defs, label)

        # Datos con cruce remapeado
        data_start_row = row_n
        remapped_rows = [_remap_cruce(r, pfx, future_years) for r in all_rows]
        row_n = _write_data_rows(ws, remapped_rows, col_defs, row_n, only_cruce=True)
        if delivery_overlay_pending:
            _overlay_section_block_label(ws, label_row, col_defs, "vigencia", delivery_date_label)
            _overlay_date_only_rows(ws, delivery_rows, col_defs, data_start_row)
            delivery_overlay_pending = False

    # ── Anchos de columna ────────────────────────────────────────────────────
    for col_idx, (sec, key, *_) in enumerate(col_defs, 1):
        if _is_sep(sec):
            ws.column_dimensions[get_column_letter(col_idx)].width = 2
        elif key == "periodo":
            ws.column_dimensions[get_column_letter(col_idx)].width = 14
        elif key in ("nro_lineas", "nro_contratos", "grupos_compra", "monto_total_usd", "yr_leq2025",
                     "yr_sin_fecha", "d_leq2025", "flag_sin", "framework_contracts",
                     "po_with_framework", "po_without_framework", "pct_po_with_framework"):
            ws.column_dimensions[get_column_letter(col_idx)].width = 14
        else:
            ws.column_dimensions[get_column_letter(col_idx)].width = 11

    ws.freeze_panes = "B4"

    if include_framework_section:
        _write_framework_detail_sheet(wb, df, framework_contract_col)

    suffix = f"-{file_suffix}" if file_suffix else ""
    path = output_dir / f"reporte_estadistico_{operation}{suffix}.xlsx"
    wb.save(str(path))
    return path
