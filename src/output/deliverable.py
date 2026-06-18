from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_INTERNAL_COLS: frozenset[str] = frozenset({
    "exclusion_reason",
    "exclusion_severity",
    "rescue_reason",
})

_HDR_DARK = "1E3A5F"
_HDR_WHITE = "FFFFFF"


def _hdr_font() -> Font:
    return Font(bold=True, color=_HDR_WHITE)


def _hdr_fill() -> PatternFill:
    return PatternFill(start_color=_HDR_DARK, end_color=_HDR_DARK, fill_type="solid")


def _to_rows(df: pd.DataFrame) -> list[list]:
    display = df.copy()
    for col in display.columns:
        if pd.api.types.is_datetime64_any_dtype(display[col]):
            display[col] = display[col].apply(lambda x: x.date() if pd.notna(x) else None)
    display = display.astype(object).where(pd.notnull(display), other=None)
    return display.values.tolist()


def _date_col_indices(df: pd.DataFrame) -> list[int]:
    return [
        i + 1
        for i, col in enumerate(df.columns)
        if pd.api.types.is_datetime64_any_dtype(df[col])
    ]


def _style_sheet(ws, df: pd.DataFrame) -> None:
    """Dark frozen header + autofilter + DD-MM-YYYY dates + autosize, sobre ws."""
    for col_idx, header in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _hdr_font()
        cell.fill = _hdr_fill()
        cell.alignment = Alignment(wrap_text=False)

    ws.freeze_panes = "A2"

    for row_data in _to_rows(df):
        ws.append(row_data)

    if not df.empty:
        ws.auto_filter.ref = ws.dimensions

    for col_idx in _date_col_indices(df):
        for cell_row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            if cell_row[0].value is not None:
                cell_row[0].number_format = "DD-MM-YYYY"

    for col_cells in ws.columns:
        width = 8
        for cell in col_cells:
            if cell.value:
                width = max(width, min(len(str(cell.value)) + 2, 42))
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width


def append_sheet_to_deliverable(
    df: pd.DataFrame,
    output_path: Path,
    sheet_name: str,
    *,
    strip_internal: bool = True,
) -> Path:
    """Anexa (o reemplaza) una hoja a un entregable existente, con el mismo estilo.

    Si el archivo no existe, lo crea con esa sola hoja.
    """
    if strip_internal:
        keep = [c for c in df.columns if c not in _INTERNAL_COLS and not c.startswith("_")]
        df = df[keep].copy()

    if output_path.exists():
        wb = load_workbook(output_path)
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
    else:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

    _style_sheet(ws, df)
    wb.save(output_path)
    return output_path


def export_deliverable(
    df: pd.DataFrame,
    output_path: Path,
    sheet_name: str,
    *,
    strip_internal: bool = True,
) -> Path:
    """Write a single-sheet production Excel deliverable.

    Dark-blue frozen header, autofilter, DD-MM-YYYY native Excel dates.
    Internal pipeline columns (exclusion_reason, …) and underscore-prefixed
    bookkeeping columns are stripped when strip_internal=True.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if strip_internal:
        keep = [c for c in df.columns if c not in _INTERNAL_COLS and not c.startswith("_")]
        df = df[keep].copy()

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    for col_idx, header in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _hdr_font()
        cell.fill = _hdr_fill()
        cell.alignment = Alignment(wrap_text=False)

    ws.freeze_panes = "A2"

    for row_data in _to_rows(df):
        ws.append(row_data)

    if not df.empty:
        ws.auto_filter.ref = ws.dimensions

    for col_idx in _date_col_indices(df):
        for cell_row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            if cell_row[0].value is not None:
                cell_row[0].number_format = "DD-MM-YYYY"

    for col_cells in ws.columns:
        width = 8
        for cell in col_cells:
            if cell.value:
                width = max(width, min(len(str(cell.value)) + 2, 42))
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width

    wb.save(output_path)
    return output_path
