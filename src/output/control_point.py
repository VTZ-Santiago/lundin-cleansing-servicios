import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.lineage.records import IssueRecord, StageManifest
from src.lineage.report import LineageReport
from src.profiling.profiler import ProfilingResult
from src.schema.canonical import CANONICAL_NAMES
from src.utils.date_logic import effective_validity_dates

# --- Colour palette ---
_HDR_DARK = "1E3A5F"
_HDR_WHITE = "FFFFFF"
_MAPPED_BG = "C8E6C9"    # green
_UNMAPPED_BG = "FFCDD2"  # red
_INJECTED_BG = "BBDEFB"  # blue
_ERROR_BG = "FFCDD2"
_WARN_BG = "FFF9C4"


def _hdr_font() -> Font:
    return Font(bold=True, color=_HDR_WHITE)


def _hdr_fill() -> PatternFill:
    return PatternFill(start_color=_HDR_DARK, end_color=_HDR_DARK, fill_type="solid")


def _write_row(ws, row_idx: int, values: list) -> None:
    for col_idx, v in enumerate(values, 1):
        ws.cell(row=row_idx, column=col_idx, value=v)


def _write_header_row(ws, headers: list[str], row: int = 1) -> None:
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = _hdr_font()
        cell.fill = _hdr_fill()
        cell.alignment = Alignment(wrap_text=False)


def _autofit(ws, max_width: int = 42) -> None:
    for col_cells in ws.columns:
        width = 8
        for cell in col_cells:
            if cell.value:
                width = max(width, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width


def _df_to_rows(df: pd.DataFrame) -> list[list]:
    """Convert DataFrame to rows safe for openpyxl. Datetime cols → Python date objects (native Excel dates)."""
    display = df.copy()
    for col in display.columns:
        if pd.api.types.is_datetime64_any_dtype(display[col]):
            display[col] = display[col].apply(lambda x: x.date() if pd.notna(x) else None)
    display = display.astype(object).where(pd.notnull(display), other=None)
    return display.values.tolist()


def _date_col_indices(df: pd.DataFrame) -> list[int]:
    """Return 1-based column indices for datetime columns (for openpyxl DD-MM-YYYY format)."""
    return [
        i + 1
        for i, col in enumerate(df.columns)
        if pd.api.types.is_datetime64_any_dtype(df[col])
    ]


# --- Analysis table helpers ---

def _extract_period(source_file: str) -> str:
    """'CONTRATO 2013-2017.xlsx' → '2013_2017'"""
    stem = Path(source_file).stem
    m = re.search(r"\d{4}[-_]\d{4}", stem)
    if m:
        return m.group().replace("-", "_")
    return stem


def _build_analysis_rows(df: pd.DataFrame) -> tuple[list[dict], list[int]]:
    """Build one summary dict per source file plus sorted year list."""
    if "_source_file" not in df.columns:
        return [], []

    years: set[int] = set()
    if "validity_end" in df.columns or "delivery_date" in df.columns:
        parsed = effective_validity_dates(df)
        years = {int(y) for y in parsed.dropna().dt.year.unique()}

    sorted_years = sorted(years)
    rows: list[dict] = []

    for src_file, subset in df.groupby("_source_file", sort=True):
        row: dict = {
            "periodo": _extract_period(str(src_file)),
            "nro_lineas": len(subset),
            "nro_contratos": int(subset["purchase_document"].nunique()) if "purchase_document" in subset.columns else 0,
        }

        if "estimated_value" in subset.columns:
            amount = pd.to_numeric(subset["estimated_value"], errors="coerce")
            row["monto_total_usd"] = round(float(amount.sum(skipna=True)), 2)
            row["monto_promedio_usd"] = round(float(amount.mean(skipna=True)), 2) if amount.notna().any() else 0.0
        else:
            row["monto_total_usd"] = 0.0
            row["monto_promedio_usd"] = 0.0

        if "purchase_group" in subset.columns:
            purchase_group = subset["purchase_group"].fillna("SIN_GRUPO").astype(str).str.strip()
            purchase_group = purchase_group.where(purchase_group != "", other="SIN_GRUPO")
            group_counts = purchase_group.value_counts()
            row["grupos_compra"] = int(purchase_group.nunique())
            row["grupo_compra_top"] = (
                f"{group_counts.index[0]} ({int(group_counts.iloc[0])})"
                if not group_counts.empty else ""
            )
        else:
            row["grupos_compra"] = 0
            row["grupo_compra_top"] = ""

        # position_type by contract: dominant (most frequent) type per purchase_document
        if "position_type" in subset.columns and "purchase_document" in subset.columns:
            pt = subset[["purchase_document", "position_type"]].copy()
            pt["position_type"] = pt["position_type"].fillna("").astype(str).str.strip()
            dominant_pt = pt.groupby("purchase_document")["position_type"].agg(
                lambda s: s.mode().iloc[0] if not s.empty else ""
            )
            vc = dominant_pt.value_counts()
            for v in ("C", "D", "V"):
                row[f"pos_{v}"] = int(vc.get(v, 0))
            row["pos_VACIO"] = int(vc.get("", 0))
        else:
            for v in ("C", "D", "V", "VACIO"):
                row[f"pos_{v}"] = 0

        # deletion_flag by contract: dominant flag per purchase_document
        if "deletion_flag" in subset.columns and "purchase_document" in subset.columns:
            fl = subset[["purchase_document", "deletion_flag"]].copy()
            fl["deletion_flag"] = fl["deletion_flag"].fillna("").astype(str).str.strip()
            dominant_fl = fl.groupby("purchase_document")["deletion_flag"].agg(
                lambda s: s.mode().iloc[0] if not s.empty else ""
            )
            vc2 = dominant_fl.value_counts()
            for v in ("L", "S"):
                row[f"flag_{v}"] = int(vc2.get(v, 0))
            row["flag_SIN_FLAG"] = int(vc2.get("", 0))
        else:
            for v in ("L", "S", "SIN_FLAG"):
                row[f"flag_{v}"] = 0

        # contracts per effective validity year (unique purchase_document per year)
        if ("validity_end" in subset.columns or "delivery_date" in subset.columns) and "purchase_document" in subset.columns:
            ve = effective_validity_dates(subset)
            for yr in sorted_years:
                mask = ve.dt.year == yr
                row[f"yr_{yr}"] = int(subset.loc[mask, "purchase_document"].nunique())
        else:
            for yr in sorted_years:
                row[f"yr_{yr}"] = 0

        rows.append(row)

    return rows, sorted_years


# --- Sheet writers ---

def _write_info(wb: Workbook, cp_id: str, description: str, operation: str,
                df: pd.DataFrame, issues: list[IssueRecord], manifests: list[StageManifest],
                analysis_rows: list[dict] | None = None,
                analysis_years: list[int] | None = None,
                analysis_subject: str = "Contratos") -> None:
    ws = wb.active
    ws.title = "Info"

    canonical_cols = [c for c in df.columns if not c.startswith("_raw__")]
    error_count = sum(1 for i in issues if i.severity == "ERROR")
    warn_count = sum(1 for i in issues if i.severity == "WARNING")
    stage_ids = ", ".join(m.stage_id for m in manifests)

    rows = [
        ("Control Point", cp_id),
        ("Descripción", description),
        ("Operación", operation),
        ("Generado el", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("Filas en master", len(df)),
        ("Columnas canónicas", len(canonical_cols)),
        ("Errores", error_count),
        ("Advertencias", warn_count),
        ("", ""),
        ("Etapas completadas", stage_ids),
    ]

    label_font = Font(bold=True)
    for row_idx, (k, v) in enumerate(rows, 1):
        ws.cell(row=row_idx, column=1, value=k).font = label_font
        ws.cell(row=row_idx, column=2, value=v)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60

    if analysis_rows:
        offset = len(rows) + 3  # blank rows before table
        years = analysis_years or []

        title_cell = ws.cell(row=offset, column=1,
                             value=f"Análisis de Datos de {analysis_subject} {operation}")
        title_cell.font = Font(bold=True, size=12)

        headers = (
            ["Periodo", "Nro de Linea", "Nro de Cttos",
             "Monto en USD", "Monto prom. USD", "Grupos de Compras", "Grupo de Compras top",
             "C", "D", "V", "Vacío",
             "L", "S", "Sin Flag"]
            + [str(y) for y in years]
        )
        _write_header_row(ws, headers, row=offset + 1)

        for i, ar in enumerate(analysis_rows, offset + 2):
            values = [
                ar.get("periodo", ""),
                ar.get("nro_lineas", 0),
                ar.get("nro_contratos", 0),
                ar.get("monto_total_usd", 0.0),
                ar.get("monto_promedio_usd", 0.0),
                ar.get("grupos_compra", 0),
                ar.get("grupo_compra_top", ""),
                ar.get("pos_C", 0),
                ar.get("pos_D", 0),
                ar.get("pos_V", 0),
                ar.get("pos_VACIO", 0),
                ar.get("flag_L", 0),
                ar.get("flag_S", 0),
                ar.get("flag_SIN_FLAG", 0),
            ] + [ar.get(f"yr_{y}", 0) for y in years]
            _write_row(ws, i, values)

        # Autofit analysis columns
        for col_cells in ws.iter_cols(min_row=offset, max_row=offset + 1 + len(analysis_rows)):
            width = 8
            for cell in col_cells:
                if cell.value:
                    width = max(width, min(len(str(cell.value)) + 2, 20))
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width


def _write_master(wb: Workbook, df: pd.DataFrame, master_columns: list[str] | None = None) -> None:
    ws = wb.create_sheet("Master")

    if master_columns is not None:
        show_cols = [column for column in master_columns if column in df.columns]
        df_display = df[show_cols].copy()
        _write_header_row(ws, show_cols)
        ws.freeze_panes = "A2"

        for row_data in _df_to_rows(df_display):
            ws.append(row_data)

        for col_idx in _date_col_indices(df_display):
            for cell_row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                if cell_row[0].value is not None:
                    cell_row[0].number_format = "DD-MM-YYYY"

        _autofit(ws)
        return

    # Columns to show: canonical order + annotations if present
    show_cols = [c for c in CANONICAL_NAMES if c in df.columns]
    for aux in ("exclusion_reason", "exclusion_severity", "rescue_reason"):
        if aux in df.columns:
            show_cols.append(aux)
    mark_cols = sorted(c for c in df.columns if c.startswith("mark_"))
    show_cols.extend(mark_cols)
    material_cols = sorted(
        c for c in df.columns
        if (c == "material_key" or c.startswith("material_")) and c not in show_cols
    )
    show_cols.extend(material_cols)

    df_display = df[show_cols].copy()
    _write_header_row(ws, show_cols)
    ws.freeze_panes = "A2"

    for row_data in _df_to_rows(df_display):
        ws.append(row_data)

    for col_idx in _date_col_indices(df_display):
        for cell_row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            if cell_row[0].value is not None:
                cell_row[0].number_format = "DD-MM-YYYY"

    _autofit(ws)


def _write_field_map(wb: Workbook, lineage: LineageReport) -> None:
    ws = wb.create_sheet("Field Map")
    headers = ["source_file", "raw_name", "canonical_name", "mapping_status",
               "null_count", "null_pct", "unique_count", "sample_values"]
    _write_header_row(ws, headers)

    status_fills = {
        "MAPPED":   PatternFill(start_color=_MAPPED_BG,   end_color=_MAPPED_BG,   fill_type="solid"),
        "UNMAPPED": PatternFill(start_color=_UNMAPPED_BG, end_color=_UNMAPPED_BG, fill_type="solid"),
        "INJECTED": PatternFill(start_color=_INJECTED_BG, end_color=_INJECTED_BG, fill_type="solid"),
    }

    for row_idx, record in enumerate(lineage.records, 2):
        d = record.to_dict()
        values = [d[h] for h in headers]
        _write_row(ws, row_idx, values)
        fill = status_fills.get(record.mapping_status)
        if fill:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    ws.freeze_panes = "A2"
    _autofit(ws)


def _write_issues(wb: Workbook, issues: list[IssueRecord]) -> None:
    ws = wb.create_sheet("Issues")
    if not issues:
        ws.cell(row=1, column=1, value="No issues recorded at this control point.")
        return

    headers = ["stage", "severity", "code", "message", "detail", "row_count"]
    _write_header_row(ws, headers)

    error_fill = PatternFill(start_color=_ERROR_BG, end_color=_ERROR_BG, fill_type="solid")
    warn_fill  = PatternFill(start_color=_WARN_BG,  end_color=_WARN_BG,  fill_type="solid")

    for row_idx, issue in enumerate(issues, 2):
        d = issue.to_dict()
        _write_row(ws, row_idx, [d[h] for h in headers])
        fill = error_fill if issue.severity == "ERROR" else (warn_fill if issue.severity == "WARNING" else None)
        if fill:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    ws.freeze_panes = "A2"
    _autofit(ws)


def _write_stages(wb: Workbook, manifests: list[StageManifest]) -> None:
    ws = wb.create_sheet("Stages")
    headers = ["stage_id", "label", "started_at", "completed_at", "duration_s",
               "rows_in", "rows_out", "columns_in", "columns_out", "notes"]
    _write_header_row(ws, headers)
    for row_idx, m in enumerate(manifests, 2):
        d = m.to_dict()
        _write_row(ws, row_idx, [d[h] for h in headers])
    ws.freeze_panes = "A2"
    _autofit(ws)


def _write_profiling(wb: Workbook, profiling: ProfilingResult) -> None:
    ws = wb.create_sheet("Profiling")
    headers = ["canonical_name", "column_group", "is_critical", "completeness_pct",
               "null_count", "unique_count", "min_val", "max_val", "mean_val", "sample_values"]
    _write_header_row(ws, headers)

    profiles_sorted = sorted(
        profiling.column_profiles,
        key=lambda p: (p.column_group, p.canonical_name)
    )
    for row_idx, p in enumerate(profiles_sorted, 2):
        d = p.to_dict()
        _write_row(ws, row_idx, [d[h] for h in headers])

    # Color scale on completeness_pct column (col D = index 4)
    if profiles_sorted:
        comp_col = "D"
        last_row = len(profiles_sorted) + 1
        ws.conditional_formatting.add(
            f"{comp_col}2:{comp_col}{last_row}",
            ColorScaleRule(
                start_type="num", start_value=0,   start_color="FF4444",
                mid_type="num",   mid_value=70,    mid_color="FFFF44",
                end_type="num",   end_value=100,   end_color="44BB44",
            ),
        )

    offset = len(profiles_sorted) + 4

    # Effective validity date by year as a summary block below the main table
    if profiling.effective_validity_by_year:
        ws.cell(row=offset, column=1, value="Distribución fecha de vigencia efectiva por año").font = Font(bold=True)
        ws.cell(row=offset, column=2, value="Contratos")
        for i, (yr, cnt) in enumerate(sorted(profiling.effective_validity_by_year.items()), 1):
            ws.cell(row=offset + i, column=1, value=yr)
            ws.cell(row=offset + i, column=2, value=cnt)
        offset += len(profiling.effective_validity_by_year) + 3

    # Source file counts
    if profiling.source_file_counts:
        ws.cell(row=offset, column=1, value="Filas por archivo de origen").font = Font(bold=True)
        ws.cell(row=offset, column=2, value="Filas")
        for i, (fname, cnt) in enumerate(profiling.source_file_counts.items(), 1):
            ws.cell(row=offset + i, column=1, value=fname)
            ws.cell(row=offset + i, column=2, value=cnt)
        offset += len(profiling.source_file_counts) + 3

    has_estimated_value = any(p.canonical_name == "estimated_value" for p in profiles_sorted)
    if has_estimated_value:
        ws.cell(row=offset, column=1, value="Resumen monto del contrato").font = Font(bold=True)
        ws.cell(row=offset + 1, column=1, value="Monto total USD")
        amount_cell = ws.cell(row=offset + 1, column=2, value=profiling.estimated_value_total)
        amount_cell.number_format = '#,##0.00'
        offset += 4

    if profiling.purchase_group_distribution:
        ws.cell(row=offset, column=1, value="Resumen por grupo de compras").font = Font(bold=True)
        _write_header_row(ws, ["Grupo de Compras", "Líneas", "Contratos", "Monto en USD"], row=offset + 1)
        groups = sorted(
            profiling.purchase_group_distribution,
            key=lambda group: (
                profiling.estimated_value_by_group.get(group, 0.0),
                profiling.purchase_group_distribution.get(group, 0),
            ),
            reverse=True,
        )
        for i, group in enumerate(groups, offset + 2):
            ws.cell(row=i, column=1, value=group)
            ws.cell(row=i, column=2, value=profiling.purchase_group_distribution.get(group, 0))
            ws.cell(row=i, column=3, value=profiling.purchase_document_count_by_group.get(group, 0))
            amount_cell = ws.cell(row=i, column=4, value=profiling.estimated_value_by_group.get(group, 0.0))
            amount_cell.number_format = '#,##0.00'

    ws.freeze_panes = "A2"
    _autofit(ws)


# --- Public API ---

def export_control_point(
    cp_id: str,
    description: str,
    operation: str,
    df: pd.DataFrame,
    lineage: LineageReport,
    profiling: ProfilingResult,
    manifests: list[StageManifest],
    issues: list[IssueRecord],
    output_dir: Path,
    include_analysis: bool = False,
    analysis_subject: str = "Contratos",
    master_columns: list[str] | None = None,
) -> Path:
    path = output_dir / f"{cp_id}_{operation}.xlsx"
    wb = Workbook()

    analysis_rows, analysis_years = (
        _build_analysis_rows(df) if include_analysis else (None, None)
    )
    _write_info(wb, cp_id, description, operation, df, issues, manifests,
                analysis_rows=analysis_rows, analysis_years=analysis_years,
                analysis_subject=analysis_subject)
    _write_master(wb, df, master_columns=master_columns)
    _write_field_map(wb, lineage)
    _write_issues(wb, issues)
    _write_stages(wb, manifests)
    _write_profiling(wb, profiling)

    wb.save(str(path))
    return path
