"""Generate the material cross-reference Excel report."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Colour palette (matches control_point.py)
_HDR_DARK = "1E3A5F"
_HDR_WHITE = "FFFFFF"
_OK_BG = "C8E6C9"       # green – found
_MISS_BG = "FFCDD2"     # red   – not found
_WARN_BG = "FFF9C4"     # yellow – partial
_SERVICE_BG = "BBDEFB"  # blue  – service material


def _hdr_font() -> Font:
    return Font(bold=True, color=_HDR_WHITE, size=10)


def _hdr_fill() -> PatternFill:
    return PatternFill(start_color=_HDR_DARK, end_color=_HDR_DARK, fill_type="solid")


def _cell_fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _write_header(ws, headers: list[str], row: int = 1) -> None:
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = _hdr_font()
        cell.fill = _hdr_fill()
        cell.alignment = Alignment(wrap_text=False)


def _autofit(ws, max_width: int = 45) -> None:
    for col_cells in ws.columns:
        width = 8
        for cell in col_cells:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width


def _df_to_sheet(ws, df: pd.DataFrame, start_row: int = 2) -> None:
    """Write DataFrame rows to worksheet starting at start_row (headers already written)."""
    display = df.copy()
    for col in display.columns:
        if pd.api.types.is_datetime64_any_dtype(display[col]):
            display[col] = display[col].dt.strftime("%Y-%m-%d").where(display[col].notna(), other=None)
    display = display.astype(object).where(pd.notnull(display), other=None)

    for r_idx, row_vals in enumerate(display.values.tolist(), start_row):
        for c_idx, val in enumerate(row_vals, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)


def _build_resumen_sheet(ws, enriched_df: pd.DataFrame, domains: list[str]) -> None:
    """Sheet 1: KPI summary."""
    ws.title = "Resumen"
    ws.freeze_panes = "B2"

    sections: list[tuple[str, object]] = [("INDICADOR", "VALOR")]

    total = len(enriched_df)
    in_master = enriched_df["in_master"].sum()
    has_consumo = enriched_df["has_consumption"].sum() if "has_consumption" in enriched_df.columns else 0
    is_service = enriched_df["is_service"].sum()

    sections += [
        ("", ""),
        ("=== COBERTURA GENERAL ===", ""),
        ("Total materiales unicos", total),
        ("Encontrados en maestro", int(in_master)),
        ("No encontrados en maestro", total - int(in_master)),
        ("% cobertura maestro", f"{in_master/total*100:.1f}%" if total else "N/A"),
        ("Con datos de consumo historico", int(has_consumo)),
        ("% cobertura consumo", f"{has_consumo/total*100:.1f}%" if total else "N/A"),
        ("", ""),
    ]

    # Breakdown by operation × domain
    operations = enriched_df["operation"].unique().tolist() if "operation" in enriched_df.columns else [None]
    for op in operations:
        sub_op = enriched_df[enriched_df["operation"] == op] if op else enriched_df
        op_label = op if op else "SIN OPERACION"
        sections.append((f"=== OPERACION: {op_label} ===", ""))
        for domain in domains:
            sub = sub_op[sub_op["domain"] == domain] if "domain" in sub_op.columns else sub_op
            n = len(sub)
            if n == 0:
                continue
            n_m = sub["in_master"].sum()
            sections += [
                (f"  Dominio: {domain}", ""),
                (f"    Materiales unicos", n),
                (f"    En maestro", int(n_m)),
                (f"    % cobertura maestro", f"{n_m/n*100:.1f}%"),
            ]
        sections.append(("", ""))

    sections += [
        ("=== MATERIALES DE SERVICIO (is_service) ===", ""),
        ("Total materiales de servicio", int(is_service)),
        ("Servicios en maestro", int(enriched_df[enriched_df["is_service"]]["in_master"].sum())),
        ("% servicios en maestro", f"{enriched_df[enriched_df['is_service']]['in_master'].sum()/is_service*100:.1f}%" if is_service else "N/A"),
        ("  Tipo posicion D", int(enriched_df["is_service_d"].sum())),
        ("  Imputacion K (centro costo)", int(enriched_df["is_direct_cost"].sum())),
        ("", ""),
        ("=== DISTRIBUCION POR TIPO MATERIAL ===", ""),
    ]

    if "tipo_material" in enriched_df.columns:
        tipo_dist = (
            enriched_df.dropna(subset=["tipo_material"])
            .groupby("tipo_material")
            .size()
            .sort_values(ascending=False)
            .head(20)
        )
        for tipo, cnt in tipo_dist.items():
            sections.append((f"  {tipo}", int(cnt)))

    _write_header(ws, ["INDICADOR", "VALOR"])
    for r_idx, (label, value) in enumerate(sections, 2):
        ws.cell(row=r_idx, column=1, value=label)
        ws.cell(row=r_idx, column=2, value=value)
        # Style section headers
        if str(label).startswith("==="):
            for c in (1, 2):
                ws.cell(row=r_idx, column=c).font = Font(bold=True, size=10)
                ws.cell(row=r_idx, column=c).fill = _cell_fill("E3F2FD")

    _autofit(ws)


def _ordered_enriched_cols(df: pd.DataFrame) -> list[str]:
    """Return column order: key info first, then master info, then consumption."""
    priority = [
        "material_key", "operation", "domain", "descripcion_efectiva", "descripcion_fuente",
        "tipos_posicion", "tipos_imputacion", "is_service", "is_service_d", "is_direct_cost",
        "n_registros", "n_documentos", "grupo_compras", "centro_fuente",
        "in_master", "has_consumption",
        # master cols
        "descripcion", "tipo_material", "centro_master", "status_master",
        "obs_1", "obs_2", "obs_3", "grupo_articulos", "segmento",
        "ind_rotacion_sugerido", "ctd_consumida_24m", "frecuencia_24m_master", "ctd_x_vez_master",
        "precio_medio_usd", "uom_base",
        # consumption cols
        "total_consumo_2024", "total_consumo_2025", "total_consumo_2026",
        "frecuencia_consumo_24m", "ctd_x_vez_consumo", "ind_rotacion_consumo",
    ]
    present = [c for c in priority if c in df.columns]
    extras = [c for c in df.columns if c not in present]
    return present + extras


def generate_report(
    enriched_df: pd.DataFrame,
    output_path: Path | str,
    domains: list[str],
) -> Path:
    """Generate the Excel cross-reference report. Returns the path of the written file."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    # --- Sheet 1: Resumen ---
    ws_res = wb.create_sheet("Resumen")
    _build_resumen_sheet(ws_res, enriched_df, domains)

    # --- Sheet 2: Materiales (all enriched) ---
    ws_mat = wb.create_sheet("Materiales")
    ordered_cols = _ordered_enriched_cols(enriched_df)
    display_df = enriched_df[ordered_cols]
    _write_header(ws_mat, ordered_cols)
    _df_to_sheet(ws_mat, display_df)
    # Color rows by in_master status
    for r_idx in range(2, len(display_df) + 2):
        in_m = display_df.iloc[r_idx - 2].get("in_master", False)
        is_srv = display_df.iloc[r_idx - 2].get("is_service", False)
        bg = _OK_BG if in_m else _MISS_BG
        if is_srv and in_m:
            bg = _SERVICE_BG
        ws_mat.cell(row=r_idx, column=1).fill = _cell_fill(bg)  # color only key column
    ws_mat.freeze_panes = "A2"
    _autofit(ws_mat)

    # --- Sheet 3: No Encontrados ---
    not_found = enriched_df[~enriched_df["in_master"]][ordered_cols].copy()
    ws_nf = wb.create_sheet("No Encontrados")
    _write_header(ws_nf, ordered_cols)
    _df_to_sheet(ws_nf, not_found)
    ws_nf.freeze_panes = "A2"
    _autofit(ws_nf)

    # --- Sheet 4: Servicios ---
    services = enriched_df[enriched_df["is_service"]][ordered_cols].copy()
    ws_srv = wb.create_sheet("Servicios")
    _write_header(ws_srv, ordered_cols)
    _df_to_sheet(ws_srv, services)
    ws_srv.freeze_panes = "A2"
    _autofit(ws_srv)

    wb.save(output_path)
    return output_path
