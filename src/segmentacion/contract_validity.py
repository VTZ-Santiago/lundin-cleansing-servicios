"""Cruce de vigencia del universo D contra los reportes de CONTRATOS VIGENTES.

Fuentes por operación:
  CCMC — inputs/CONTRATOS VIGENTES CCMC <MES> <AA>.xlsx
    - Hoja "Report Cttos Vigentes": contratos marco (46*) con FIN CTTO SAP (fallback Fin (Orden)).
    - Hoja "Report Pedidos": documentos de compra (45*, tipos Ctto y OST) con columna Fin.
  MLCC — inputs/Reporte Contratos Vigentes <Mes>-<AA>.xlsx
    - Hoja "VIGENTES": catálogo curado que mezcla ambas operaciones; se filtra Sitio=CAS (MLCC).
      Clave "No SAP" (marcos 46* y contratos-documento 51/52/53/54*), fecha "Fin Ctto".
    - Fallback de compatibilidad: inputs/MLCC/CONTRATOS VIGENTES MLCC <MES> <AA>.xlsx, hoja
      "CONSOL" (clave "Documento compras", fecha "Fin período validez").

Clave de cruce por fila OC: outline_contract si existe; si no, purchase_document.
Las filas sin cruce quedan con fecha NaT — E01 las trata como sin contrato vigente.
"""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

END_DATE_COLUMN = "contract_end_date"
MATCH_SOURCE_COLUMN = "contract_match_source"

MATCH_BY_CONTRACT = "CONTRATO_MARCO"
MATCH_BY_DOCUMENT = "DOCUMENTO_COMPRA"
NO_MATCH = "SIN_CRUCE"

_HEADER_SCAN_ROWS = 8


@dataclass
class VigentesIndex:
    operation: str
    source_file: str
    end_dates: dict[str, pd.Timestamp] = field(default_factory=dict)
    sheet_counts: dict[str, int] = field(default_factory=dict)

    @property
    def total_keys(self) -> int:
        return len(self.end_dates)


def _normalize_key(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not text or text.upper() in {"NAN", "NONE", "NAT", "0"}:
        return ""
    return text


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.split()).strip().lower()


def _find_sheet(workbook, name: str):
    target = _normalize_header(name)
    for sheet_name in workbook.sheetnames:
        if _normalize_header(sheet_name) == target:
            return workbook[sheet_name]
    for sheet_name in workbook.sheetnames:
        if target in _normalize_header(sheet_name):
            return workbook[sheet_name]
    raise KeyError(f"No se encontró la hoja '{name}' en el reporte de vigentes")


def _load_sheet_dates(
    ws,
    key_headers: list[str],
    date_headers: list[str],
    filter_headers: list[str] | None = None,
    filter_values: set[str] | None = None,
) -> dict[str, pd.Timestamp]:
    """Map key → max end date from a sheet. Header row is auto-detected.

    key_headers / date_headers are ordered by preference; the first non-empty
    value per row wins. When filter_headers/filter_values are given, only rows
    whose normalized value in the filter column is in filter_values are kept.
    """
    rows = ws.iter_rows(values_only=True)

    key_cols: list[int] = []
    date_cols: list[int] = []
    filter_cols: list[int] = []
    for _ in range(_HEADER_SCAN_ROWS):
        row = next(rows, None)
        if row is None:
            break
        headers = {_normalize_header(v): idx for idx, v in enumerate(row) if v is not None}
        key_cols = [headers[_normalize_header(h)] for h in key_headers if _normalize_header(h) in headers]
        date_cols = [headers[_normalize_header(h)] for h in date_headers if _normalize_header(h) in headers]
        if key_cols and date_cols:
            if filter_headers:
                filter_cols = [headers[_normalize_header(h)] for h in filter_headers if _normalize_header(h) in headers]
            break
        key_cols = []
        date_cols = []
    if not key_cols or not date_cols:
        raise ValueError(f"No se encontraron encabezados {key_headers}/{date_headers} en la hoja '{ws.title}'")
    if filter_headers and not filter_cols:
        raise ValueError(f"No se encontró el encabezado de filtro {filter_headers} en la hoja '{ws.title}'")

    wanted = {str(v).strip().upper() for v in filter_values} if filter_values else None

    end_dates: dict[str, pd.Timestamp] = {}
    for row in rows:
        if wanted is not None:
            value = ""
            for idx in filter_cols:
                if idx < len(row) and row[idx] is not None:
                    value = str(row[idx]).strip().upper()
                    break
            if value not in wanted:
                continue
        key = ""
        for idx in key_cols:
            if idx < len(row):
                key = _normalize_key(row[idx])
                if key:
                    break
        if not key:
            continue
        end = pd.NaT
        for idx in date_cols:
            if idx < len(row):
                end = pd.to_datetime(row[idx], errors="coerce")
                if pd.notna(end):
                    break
        previous = end_dates.get(key)
        if previous is None or (pd.notna(end) and (pd.isna(previous) or end > previous)):
            end_dates[key] = end
    return end_dates


def _latest_file(directory: Path, pattern: str) -> Path:
    candidates = [p for p in directory.glob(pattern) if not p.name.startswith("~$")]
    if not candidates:
        raise FileNotFoundError(f"No se encontró '{pattern}' en {directory}")
    return max(candidates, key=lambda p: p.stat().st_mtime)


def load_vigentes_index(operation: str, inputs_dir: Path) -> VigentesIndex:
    operation = operation.upper()
    if operation == "CCMC":
        pattern = "CONTRATOS VIGENTES CCMC*.xlsx"
        if list((inputs_dir / "CCMC").glob(pattern)):
            path = _latest_file(inputs_dir / "CCMC", pattern)
        else:
            path = _latest_file(inputs_dir, pattern)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            marcos = _load_sheet_dates(
                _find_sheet(workbook, "Report Cttos Vigentes"),
                key_headers=["Contrato marco2", "Contrato marco"],
                date_headers=["FIN CTTO SAP", "Fin (Orden)"],
            )
            pedidos = _load_sheet_dates(
                _find_sheet(workbook, "Report Pedidos"),
                key_headers=["Documento compras"],
                date_headers=["Fin"],
            )
        finally:
            workbook.close()
        index = VigentesIndex(operation=operation, source_file=path.name)
        index.end_dates.update(pedidos)
        index.end_dates.update(marcos)  # marcos prevalecen ante colisión de clave
        index.sheet_counts = {"Report Cttos Vigentes": len(marcos), "Report Pedidos": len(pedidos)}
        return index

    if operation == "MLCC":
        if list(inputs_dir.glob("Reporte Contratos Vigentes*.xlsx")):
            path = _latest_file(inputs_dir, "Reporte Contratos Vigentes*.xlsx")
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                vigentes = _load_sheet_dates(
                    _find_sheet(workbook, "VIGENTES"),
                    key_headers=["No SAP"],
                    date_headers=["Fin Ctto"],
                    filter_headers=["Sitio"],
                    filter_values={"CAS"},
                )
            finally:
                workbook.close()
            index = VigentesIndex(operation=operation, source_file=path.name)
            index.end_dates.update(vigentes)
            index.sheet_counts = {"VIGENTES (CAS)": len(vigentes)}
            return index

        # Fallback de compatibilidad: reporte consolidado MLCC con hoja CONSOL.
        path = _latest_file(inputs_dir / "MLCC", "CONTRATOS VIGENTES MLCC*.xlsx")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            consol = _load_sheet_dates(
                _find_sheet(workbook, "CONSOL"),
                key_headers=["Documento compras"],
                date_headers=["Fin período validez"],
            )
        finally:
            workbook.close()
        index = VigentesIndex(operation=operation, source_file=path.name)
        index.end_dates.update(consol)
        index.sheet_counts = {"CONSOL": len(consol)}
        return index

    raise ValueError(f"Operación no soportada para cruce de vigentes: {operation}")


def apply_contract_validity(df: pd.DataFrame, index: VigentesIndex) -> tuple[pd.DataFrame, dict[str, int]]:
    """Add contract_end_date / contract_match_source columns by crossing the vigentes index."""
    result = df.copy()
    marco_keys = result.get("outline_contract", pd.Series("", index=result.index)).map(_normalize_key)
    doc_keys = result.get("purchase_document", pd.Series("", index=result.index)).map(_normalize_key)
    use_marco = marco_keys != ""
    lookup = marco_keys.where(use_marco, doc_keys)

    end_dates = pd.to_datetime(lookup.map(index.end_dates), errors="coerce")
    matched = end_dates.notna()

    source = pd.Series(NO_MATCH, index=result.index, dtype="object")
    source[matched & use_marco] = MATCH_BY_CONTRACT
    source[matched & ~use_marco] = MATCH_BY_DOCUMENT

    result[END_DATE_COLUMN] = end_dates
    result[MATCH_SOURCE_COLUMN] = source

    stats = {
        "rows": len(result),
        "matched": int(matched.sum()),
        "by_contract": int((matched & use_marco).sum()),
        "by_document": int((matched & ~use_marco).sum()),
        "sin_cruce": int((~matched).sum()),
    }
    return result, stats
