"""OST vigentes sin contrato madre (45* sin 46*) que el flujo deja fuera.

El cliente (Candelaria) marca como vigentes ciertas órdenes de servicio cuyo
documento de compra es 45* y NO cuelgan de un contrato marco (46*). Como el
cruce de vigencia del flujo principal no las captura, aquí se identifican las
que NO están ya en los entregables que migran y se enriquecen con la base OC
para anexarlas como una hoja extra del entregable de contratos.

Solo aplica a CCMC y solo si existe el registro en `inputs/CCMC/` con la hoja
"OS Vigentes". No toca controles ni la lógica de migración del flujo.
"""
from __future__ import annotations

import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

REGISTRY_SHEET = "OS Vigentes"
MATCH_SOURCE = "OST_VIGENTE_SIN_MARCO"


def _norm_header(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.split()).strip().lower()


def _norm_key(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if text.upper() in {"NAN", "NONE", "NAT"}:
        return ""
    return text


def find_registry(source_dir: Path) -> Path | None:
    """Devuelve el xlsx en una carpeta fuente que contenga la hoja 'OS Vigentes' (más reciente)."""
    candidates = sorted(
        (p for p in source_dir.glob("*Vigentes*.xlsx") if not p.name.startswith("~$")),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    target = _norm_header(REGISTRY_SHEET)
    for path in candidates:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            if any(_norm_header(name) == target for name in wb.sheetnames):
                return path
        finally:
            wb.close()
    return None


def load_os_vigentes(path: Path) -> pd.DataFrame:
    """Lee la hoja OS Vigentes → columnas doc, pos, contract_end_date (Fin)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next(
            (wb[name] for name in wb.sheetnames if _norm_header(name) == _norm_header(REGISTRY_SHEET)),
            None,
        )
        if sheet is None:
            return pd.DataFrame(columns=["doc", "pos", "contract_end_date"])
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return pd.DataFrame(columns=["doc", "pos", "contract_end_date"])
        idx = {}
        for i, value in enumerate(header):
            key = _norm_header(value)
            if key and key not in idx:
                idx[key] = i
        doc_i = idx.get("documento compras")
        pos_i = idx.get("posicion")
        fin_i = idx.get("fin")
        if doc_i is None:
            return pd.DataFrame(columns=["doc", "pos", "contract_end_date"])
        records = []
        for row in rows:
            doc = _norm_key(row[doc_i]) if doc_i < len(row) else ""
            if not doc:
                continue
            pos = _norm_key(row[pos_i]) if pos_i is not None and pos_i < len(row) else ""
            fin = row[fin_i] if fin_i is not None and fin_i < len(row) else None
            records.append({"doc": doc, "pos": pos, "contract_end_date": fin})
        df = pd.DataFrame.from_records(records, columns=["doc", "pos", "contract_end_date"])
        df["contract_end_date"] = pd.to_datetime(df["contract_end_date"], errors="coerce")
        return df
    finally:
        wb.close()


def build_missing_ost_vigentes(
    master: pd.DataFrame,
    migra_documents: set[str],
    registry_path: Path,
    columns: list[str],
) -> pd.DataFrame:
    """DataFrame de las OST vigentes (45* sin marco) que NO están ya migrando.

    Grano (documento, posición) tal como el registro las lista; se enriquecen con
    la base `master` por esa clave. `migra_documents` son los documentos que ya
    migran (contratos ∪ OS) y se excluyen.
    """
    reg = load_os_vigentes(registry_path)
    if reg.empty:
        return pd.DataFrame(columns=columns)

    migra = {_norm_key(d) for d in migra_documents}
    missing = reg[~reg["doc"].isin(migra)].copy()
    if missing.empty:
        return pd.DataFrame(columns=columns)

    base = master.copy()
    base["__doc"] = base.get("purchase_document", pd.Series("", index=base.index)).map(_norm_key)
    base["__pos"] = base.get("position", pd.Series("", index=base.index)).map(_norm_key)
    # una fila por (doc,pos) de la base (evita duplicar si hubiese repetidos)
    base = base.drop_duplicates(subset=["__doc", "__pos"], keep="first")

    merged = missing.merge(
        base,
        left_on=["doc", "pos"],
        right_on=["__doc", "__pos"],
        how="left",
        suffixes=("", "_base"),
    )

    out = pd.DataFrame(index=merged.index)
    for col in columns:
        if col in merged.columns:
            out[col] = merged[col]
        else:
            out[col] = pd.NA
    # claves desde el registro (autoritativo) y marca de vigencia
    out["purchase_document"] = merged["doc"].values
    out["position"] = merged["pos"].values
    out["contract_end_date"] = pd.to_datetime(merged["contract_end_date"], errors="coerce").values
    if "contract_match_source" in out.columns:
        out["contract_match_source"] = MATCH_SOURCE
    if "migration_category" in out.columns:
        out["migration_category"] = "VIGENTE"
    if "delivery_year" in out.columns:
        out["delivery_year"] = pd.to_datetime(out["contract_end_date"], errors="coerce").dt.year.astype("Int64")
    return out[columns]
