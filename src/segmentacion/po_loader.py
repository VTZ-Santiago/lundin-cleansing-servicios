from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook

from src.lineage.records import FieldLineageRecord, StageManifest
from src.lineage.report import LineageReport
from src.ordenes_compra.schema.field_map_mlcc import MLCC_DTYPE_COERCIONS, MLCC_RAW_TO_CANONICAL


EXTRA_RAW_TO_CANONICAL = {
    "PR/SOLPED": "purchase_requisition",
    "Purchase Requisition": "purchase_requisition",
    "Purchase Requisition.1": "purchase_requisition",
    "Contrato marco": "outline_contract",
    "Contrato Marco": "outline_contract",
    "Outline Agreement": "outline_contract",
    "Planta": "plant_code",
    "Plant": "plant_code",
    "Purchasing Document": "purchase_document",
    "Item": "position",
    "Vendor/supplying plant": "vendor",
    "Vendor/Supplying Plant": "vendor",
    "Short Text": "short_text",
    "Item Category": "position_type",
    "Item Category.1": "position_type",
    "Purchasing Doc. Type": "purchase_doc_class",
    "Purch. Doc. Category": "purchase_doc_class",
    "Purchasing Group": "purchase_group",
    "Release group": "release_group",
    "Release Group": "release_group",
    "Deletion Flag": "deletion_flag",
    "Deletion Indicator": "deletion_flag",
    "Still to be delivered (qty)": "pending_delivery_qty",
    "Still to be delivered (value)": "pending_delivery_value",
    "Document Date": "document_date",
    "Delivery Date": "delivery_date",
    "Fecha de Termino": "validity_end",
    "Fecha de término": "validity_end",
    "Validity Period End": "validity_end",
    "Acct Assignment Cat.": "account_assignment_type",
}

RAW_HEADER_PRIORITY: dict[str, int] = {
    "purchasing doc. type": 20,
    "purch. doc. category": 10,
    "purchase requisition.1": 20,
    "purchase requisition": 10,
    "tipo de posicion": 20,
    "item category.1": 20,
    "item category": 10,
}

RAW_HEADER_OCCURRENCE_PRIORITY: dict[tuple[str, int], int] = {
    ("purchase requisition", 2): 20,
    ("purchase requisition", 1): 10,
    ("item category", 2): 20,
    ("item category", 1): 10,
}

WANTED_CANONICAL_COLUMNS = {
    "account_assignment_type",
    "deletion_flag",
    "delivery_date",
    "document_date",
    "material",
    "outline_contract",
    "pending_delivery_qty",
    "pending_delivery_value",
    "plant_code",
    "position",
    "position_type",
    "purchase_doc_class",
    "purchase_doc_type",
    "purchase_document",
    "purchase_group",
    "release_group",
    "purchase_requisition",
    "short_text",
    "validity_end",
    "vendor",
}


@dataclass
class SourceFileStats:
    source_file: str
    rows_read: int = 0
    rows_kept: int = 0
    missing_mapped_headers: list[str] = field(default_factory=list)
    skipped: bool = False
    skip_reason: str = ""


@dataclass
class PurchaseOrderLoadResult:
    dataframe: pd.DataFrame
    lineage: LineageReport
    manifests: list[StageManifest]
    file_stats: list[SourceFileStats]

    @property
    def rows_read(self) -> int:
        return sum(stat.rows_read for stat in self.file_stats)

    @property
    def rows_kept(self) -> int:
        return sum(stat.rows_kept for stat in self.file_stats)


def _normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value).strip())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"\s+", " ", text)
    return text.lower()


def _header_priority(header: object, occurrence: int = 1) -> int:
    normalized = _normalize_header(header)
    return RAW_HEADER_OCCURRENCE_PRIORITY.get(
        (normalized, occurrence),
        RAW_HEADER_PRIORITY.get(normalized, 10),
    )


def _mapping_by_normalized_header() -> dict[str, str]:
    mapping = dict(MLCC_RAW_TO_CANONICAL)
    mapping.update(EXTRA_RAW_TO_CANONICAL)
    return {
        _normalize_header(raw): canonical
        for raw, canonical in mapping.items()
        if canonical in WANTED_CANONICAL_COLUMNS
    }


def _excel_files(inputs_root: Path, operation: str) -> list[Path]:
    directory = inputs_root / operation.upper() / "ordenes-compra"
    if not directory.exists():
        raise FileNotFoundError(f"Input directory not found: {directory}")
    files = [
        path
        for path in list(directory.glob("*.xlsx")) + list(directory.glob("*.XLSX"))
        if not path.name.startswith("~$")
    ]
    unique = sorted({path.resolve(): path for path in files}.values(), key=lambda path: path.name.lower())
    if not unique:
        raise FileNotFoundError(f"No Excel files found in {directory}")
    return unique


def _selected_columns(headers: tuple[object, ...], mapping: dict[str, str]) -> dict[str, tuple[int, str]]:
    selected: dict[str, tuple[int, str, int]] = {}
    occurrences: dict[str, int] = {}
    for idx, raw_header in enumerate(headers):
        normalized = _normalize_header(raw_header)
        occurrences[normalized] = occurrences.get(normalized, 0) + 1
        canonical = mapping.get(normalized)
        if canonical is None:
            continue
        priority = _header_priority(raw_header, occurrences[normalized])
        current = selected.get(canonical)
        if current is None or priority > current[2]:
            selected[canonical] = (idx, str(raw_header).strip(), priority)
    return {canonical: (idx, raw) for canonical, (idx, raw, _) in selected.items()}


def _clean_text(value: object) -> object:
    if pd.isna(value):
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _coerce_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for column in result.columns:
        if column.startswith("_"):
            continue
        dtype = MLCC_DTYPE_COERCIONS.get(column, "str")
        if dtype == "date":
            result[column] = pd.to_datetime(result[column], errors="coerce")
        elif dtype == "float":
            result[column] = pd.to_numeric(result[column], errors="coerce")
        else:
            result[column] = result[column].map(_clean_text)
    return result


def _has_selected_values(row_values: tuple[object, ...], selected_columns: dict[str, tuple[int, str]]) -> bool:
    for idx, _ in selected_columns.values():
        if idx < len(row_values) and not pd.isna(row_values[idx]):
            return True
    return False


def _stream_file(path: Path, operation: str, mapping: dict[str, str]) -> tuple[pd.DataFrame, SourceFileStats, dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None:
            return pd.DataFrame(), SourceFileStats(path.name, skipped=True, skip_reason="Empty workbook"), {}

        selected_columns = _selected_columns(header_row, mapping)
        if "purchase_document" not in selected_columns or "position" not in selected_columns:
            return pd.DataFrame(), SourceFileStats(
                path.name,
                skipped=True,
                skip_reason="Missing purchase_document or position",
            ), {}

        raw_by_canonical = {canonical: raw for canonical, (_, raw) in selected_columns.items()}
        canonical_columns = [canonical for canonical, _ in sorted(selected_columns.items(), key=lambda item: item[1][0])]

        rows: list[dict[str, object]] = []
        rows_read = 0
        last_purchase_document = None
        max_selected_col = max(idx for idx, _ in selected_columns.values()) + 1
        for row_values in ws.iter_rows(min_row=2, max_col=max_selected_col, values_only=True):
            if not _has_selected_values(row_values, selected_columns):
                continue
            rows_read += 1
            row = {
                canonical: row_values[idx] if idx < len(row_values) else None
                for canonical, (idx, _) in selected_columns.items()
            }
            if pd.isna(row.get("purchase_document")):
                row["purchase_document"] = last_purchase_document
            else:
                last_purchase_document = row.get("purchase_document")
            if pd.isna(row.get("purchase_document")) and pd.isna(row.get("position")):
                continue
            rows.append(row)

        df = pd.DataFrame(rows, columns=canonical_columns)
        df["_source_file"] = path.name
        df["_operation"] = operation.upper()
        df = _coerce_dataframe(df)

        return df, SourceFileStats(path.name, rows_read=rows_read, rows_kept=len(df)), raw_by_canonical
    finally:
        wb.close()


def _build_lineage(df: pd.DataFrame, raw_names: dict[str, set[str]]) -> LineageReport:
    report = LineageReport()
    total = max(len(df), 1)
    for column in sorted(c for c in df.columns if not c.startswith("_")):
        series = df[column]
        null_count = int(series.isna().sum())
        report.add(FieldLineageRecord(
            source_file="(ordenes-compra)",
            raw_name=" | ".join(sorted(raw_names.get(column, {column}))),
            canonical_name=column,
            mapping_status="MAPPED",
            null_count=null_count,
            null_pct=round(null_count / total * 100, 2),
            unique_count=int(series.nunique(dropna=True)),
            sample_values=[str(value) for value in series.dropna().iloc[:5].tolist()],
        ))

    for column in ("_source_file", "_operation"):
        if column not in df.columns:
            continue
        report.add(FieldLineageRecord(
            source_file="(segmentacion)",
            raw_name="",
            canonical_name=column,
            mapping_status="INJECTED",
            null_count=0,
            null_pct=0.0,
            unique_count=int(df[column].nunique(dropna=True)),
            sample_values=[str(value) for value in df[column].dropna().unique()[:5].tolist()],
        ))
    return report


def load_purchase_orders(inputs_root: Path, operation: str = "MLCC") -> PurchaseOrderLoadResult:
    started = datetime.now()
    mapping = _mapping_by_normalized_header()
    frames: list[pd.DataFrame] = []
    stats: list[SourceFileStats] = []
    raw_names: dict[str, set[str]] = {}

    for path in _excel_files(inputs_root, operation):
        print(f"  Leyendo [{operation.upper()}] {path.name}", flush=True)
        file_df, file_stat, raw_by_canonical = _stream_file(path, operation, mapping)
        if not file_df.empty:
            frames.append(file_df)
        stats.append(file_stat)
        for canonical, raw in raw_by_canonical.items():
            raw_names.setdefault(canonical, set()).add(raw)

    df = pd.concat(frames, ignore_index=True, sort=False) if frames else pd.DataFrame()
    if not df.empty:
        leading = [c for c in ("purchase_document", "position", "purchase_doc_class", "position_type") if c in df.columns]
        trailing = [c for c in ("_source_file", "_operation") if c in df.columns]
        middle = [c for c in df.columns if c not in leading and c not in trailing]
        df = df[leading + middle + trailing]

    lineage = _build_lineage(df, raw_names)
    manifest = StageManifest(
        stage_id="PO_STREAM_LOAD",
        label=f"Carga streaming de ordenes de compra {operation.upper()}",
        started_at=started,
        completed_at=datetime.now(),
        rows_in=sum(stat.rows_read for stat in stats),
        rows_out=len(df),
        columns_in=0,
        columns_out=len(df.columns),
        notes=f"Archivos cargados: {', '.join(stat.source_file for stat in stats if not stat.skipped)}",
    )
    return PurchaseOrderLoadResult(df, lineage, [manifest], stats)
