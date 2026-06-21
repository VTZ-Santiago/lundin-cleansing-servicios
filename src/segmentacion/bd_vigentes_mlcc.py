"""OS vigentes del catálogo del cliente MLCC (Sitio Caserones_BD <mes>).

El cliente entrega una base (`inputs/MLCC/Sitio Caserones_BD *.xlsx`, hoja
"SAP") donde la vigencia de las órdenes de servicio vive en `Fin período
validez` (no en `delivery_date`). El flujo principal de MLCC toma las OS de la
base de suministros y mide vigencia por fecha de entrega, así que estas OS
—vigentes según el cliente— no aparecen como vigentes en el entregable.

Aquí se identifican las OS (tipo D, sin marco, clases OS) vigentes por
`Fin período validez`, con saldo pendiente y sin indicador de borrado, y se
mapean al esquema del entregable para anexarlas como una hoja extra. No toca
controles ni la lógica de migración del flujo.
"""
from __future__ import annotations

import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SHEET = "SAP"
MATCH_SOURCE = "BD_CLIENTE_FIN_VALIDEZ"
OS_CLASSES = {"ZADI", "ZAUT", "ZLVS", "ZPD1", "ZSEM", "ZSPF"}

# header SAP (normalizado) -> columna canónica del entregable
_HEADER_TO_CANONICAL = {
    "su referencia": "purchase_requisition",
    "contrato marco": "outline_contract",
    "documento compras": "purchase_document",
    "posicion": "position",
    "nombre de proveedor": "vendor",
    "texto breve": "short_text",
    "fecha documento": "document_date",
    "fin periodo validez": "validity_end",
    "por entregar (valor)": "pending_delivery_value",
    "por entregar (cantidad)": "pending_delivery_qty",
    "cl documento compras": "purchase_doc_class",
    "cl.documento compras": "purchase_doc_class",
    "grupo de compras": "purchase_group",
    "grupo de liberacion": "release_group",
    "indicador de borrado": "deletion_flag",
    "tipo de posicion": "position_type",
    "moneda": "currency",
}

_POSITION_TYPE_MAP = {"0": "", "2": "K", "3": "L", "9": "D"}


def _norm_header(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.split()).strip().lower()


def _norm_key(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return "" if text.upper() in {"NAN", "NONE", "NAT", "0"} else text


def find_bd_base(inputs_mlcc_dir: Path) -> Path | None:
    """Devuelve el `Sitio Caserones_BD*.xlsx` más reciente (o None)."""
    candidates = [
        p for p in inputs_mlcc_dir.glob("Sitio Caserones_BD*.xlsx")
        if not p.name.startswith("~$")
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def _load_sap(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next((wb[n] for n in wb.sheetnames if _norm_header(n) == _norm_header(SHEET)), None)
        if sheet is None:
            return pd.DataFrame()
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return pd.DataFrame()
        # primera ocurrencia de cada canónica
        col_for_canonical: dict[str, int] = {}
        for i, raw in enumerate(header):
            canonical = _HEADER_TO_CANONICAL.get(_norm_header(raw))
            if canonical and canonical not in col_for_canonical:
                col_for_canonical[canonical] = i
        records: list[dict] = []
        for row in rows:
            if not row or all(v is None for v in row):
                continue
            records.append({c: (row[i] if i < len(row) else None) for c, i in col_for_canonical.items()})
        return pd.DataFrame.from_records(records)
    finally:
        wb.close()


def build_os_vigentes_bd(path: Path, columns: list[str], cutoff_date: str) -> pd.DataFrame:
    """OS vigentes por `Fin período validez` (con saldo, sin borrado) → esquema entregable."""
    df = _load_sap(path)
    if df.empty:
        return pd.DataFrame(columns=columns)

    cutoff = pd.to_datetime(cutoff_date, errors="coerce")
    pos_type = df.get("position_type", pd.Series("", index=df.index)).fillna("").astype(str).str.strip()
    pos_type_canon = pos_type.map(lambda v: _POSITION_TYPE_MAP.get(v, v.upper()))
    doc_class = df.get("purchase_doc_class", pd.Series("", index=df.index)).fillna("").astype(str).str.strip().str.upper()
    marco = df.get("outline_contract", pd.Series("", index=df.index)).map(_norm_key)
    vendor = df.get("vendor", pd.Series("", index=df.index)).fillna("").astype(str).str.strip()
    short_text = df.get("short_text", pd.Series("", index=df.index)).fillna("").astype(str).str.strip()
    fin = pd.to_datetime(df.get("validity_end", pd.Series(pd.NaT, index=df.index)), errors="coerce")
    value = pd.to_numeric(df.get("pending_delivery_value", pd.Series(0, index=df.index)), errors="coerce").fillna(0)
    qty = pd.to_numeric(df.get("pending_delivery_qty", pd.Series(0, index=df.index)), errors="coerce").fillna(0)
    deletion = df.get("deletion_flag", pd.Series("", index=df.index)).fillna("").astype(str).str.strip()

    os_scope = (
        (pos_type_canon == "D")
        & (marco == "")
        & doc_class.isin(OS_CLASSES)
        & (vendor != "")
        & (short_text != "")
    )
    vigente = fin.notna() & (fin > cutoff)
    con_saldo = (value > 0) | (qty > 0)
    sin_borrado = deletion == ""
    selected = df[os_scope & vigente & con_saldo & sin_borrado].copy()
    if selected.empty:
        return pd.DataFrame(columns=columns)

    out = pd.DataFrame(index=selected.index)
    for col in columns:
        out[col] = selected[col] if col in selected.columns else pd.NA
    out["position_type"] = pos_type_canon.loc[selected.index].values
    out["document_date"] = pd.to_datetime(selected.get("document_date"), errors="coerce").values
    end = pd.to_datetime(selected["validity_end"], errors="coerce")
    out["validity_end"] = end.values
    if "contract_end_date" in out.columns:
        out["contract_end_date"] = end.values
    if "contract_match_source" in out.columns:
        out["contract_match_source"] = MATCH_SOURCE
    if "migration_category" in out.columns:
        out["migration_category"] = "VIGENTE"
    if "delivery_year" in out.columns:
        out["delivery_year"] = end.dt.year.astype("Int64").values
    return out[columns]
