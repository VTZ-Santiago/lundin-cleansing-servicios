"""Load and normalize data from ordenes-compra and ME3L sources."""
import re
import unicodedata
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import load_workbook

from src.ordenes_compra.po_consolidator import consolidate_purchase_orders


# ---------------------------------------------------------------------------
# Header / value normalization
# ---------------------------------------------------------------------------

def _normalize(s: object) -> str:
    """NFC → strip combining → lowercase → collapse whitespace."""
    text = unicodedata.normalize("NFC", str(s)).strip()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).lower()


def _norm_material(value: object) -> Optional[str]:
    """SAP material value → clean integer string, or None."""
    if value is None or value == "":
        return None
    if isinstance(value, float):
        if pd.isna(value) or value == 0.0:
            return None
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip()
    if not text or text.lower() in ("none", "nan"):
        return None
    try:
        num = float(text)
        if num == 0:
            return None
        return str(int(num)) if num.is_integer() else text
    except (ValueError, OverflowError):
        pass
    return text


def _norm_id(value: object) -> Optional[str]:
    """Document number / position → clean string, or None."""
    if value is None or value == "":
        return None
    if isinstance(value, float):
        if pd.isna(value):
            return None
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip()
    if not text or text.lower() in ("none", "nan"):
        return None
    try:
        num = float(text)
        return str(int(num)) if num.is_integer() else text
    except (ValueError, OverflowError):
        pass
    return text


# ---------------------------------------------------------------------------
# Column maps
# ---------------------------------------------------------------------------

# Normalized header → unified output column name (ME3L source)
_ME3L_COL_MAP: dict[str, str] = {
    "centro":                               "Centro",
    "organizacion compras":                 "Org. Compras",
    "contrato marco":                       "Contrato Marco",
    "pos.contrato sup.":                    "Pos. Ctto Sup.",
    "documento compras":                    "Documento compras",
    "posicion":                             "Posición",
    "material":                             "Material",
    "texto breve":                          "Texto Breve",
    "proveedor/centro suministrador":       "Proveedor",
    "cl.documento compras":                 "Cl Doc Compras",
    "tipo doc.compras":                     "Tipo Doc Compras",
    "tipo de posicion":                     "Tipo de Posición",
    "tipo de imputacion":                   "Tipo Imputación",
    "grupo de articulos":                   "Grupo Artículos",
    "indicador de borrado":                 "Ind. Borrado",
    "in.periodo validez":                   "Inicio Validez",
    "fin periodo validez":                  "Fecha Término",
    "fecha documento":                      "Fecha Documento",
    "por entregar (cantidad)":              "Por Entregar (Qty)",
    "por entregar (valor)":                 "Por Entregar (Valor)",
    "historial pedido/docu.orden entrega":  "Historial Pedido",
}

# OC consolidado column → unified output column name
_OC_COL_MAP: dict[str, str] = {
    "Planta":                   "Centro",
    "PR/SOLPED":                "PR/SOLPED",
    "Contrato Marco":           "Contrato Marco",
    "Documento compras":        "Documento compras",
    "Posición":                 "Posición",
    "Material":                 "Material",
    "Proveedor":                "Proveedor",
    "Texto Breve":              "Texto Breve",
    "Tipo de Posición":         "Tipo de Posición",
    "Cl Docto compras":         "Cl Doc Compras",
    "Grupo de liberación":      "Grupo Liberación",
    "Por entregar (cantidad)":  "Por Entregar (Qty)",
    "Por entregar (valor)":     "Por Entregar (Valor)",
    "Fecha documento":          "Fecha Documento",
    "Fecha de Entrega":         "Fecha Entrega",
    "Fecha de Termino":         "Fecha Término",
    "Indicador de borrado":     "Ind. Borrado",
}


# ---------------------------------------------------------------------------
# ME3L loader
# ---------------------------------------------------------------------------

def load_me3l(inputs_root: Path, operation: str) -> pd.DataFrame:
    """Find and load ME3L_Ctto Excel for the given operation."""
    op_dir = inputs_root / operation.upper()
    candidates = sorted(
        {p.resolve(): p
         for p in (list(op_dir.glob("ME3L_Ctto*.XLSX")) + list(op_dir.glob("ME3L_Ctto*.xlsx")))
         if not p.name.startswith("~$")}.values()
    )
    if not candidates:
        return pd.DataFrame()
    return _load_me3l_file(candidates[0], operation.upper())


def _load_me3l_file(path: Path, operation: str) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)

        header_row = next(rows_iter, None)
        if not header_row:
            return pd.DataFrame()

        # Map raw header → output column (first occurrence wins for duplicate headers)
        col_map: dict[str, int] = {}
        for idx, raw_h in enumerate(header_row):
            if raw_h is None:
                continue
            key = _normalize(raw_h)
            out = _ME3L_COL_MAP.get(key)
            if out and out not in col_map:
                col_map[out] = idx

        if "Documento compras" not in col_map or "Material" not in col_map:
            return pd.DataFrame()

        raw_records: list[dict] = []
        for row_values in rows_iter:
            rec = {
                out: (row_values[idx] if idx < len(row_values) else None)
                for out, idx in col_map.items()
            }
            raw_records.append(rec)

        if not raw_records:
            return pd.DataFrame()

        df = pd.DataFrame(raw_records)

        # Forward-fill Documento compras (SAP prints it once per contract block)
        df["Documento compras"] = df["Documento compras"].apply(_norm_id)
        df["Documento compras"] = df["Documento compras"].ffill()

        # Normalize material and drop rows without both keys
        df["Material"] = df["Material"].apply(_norm_material)
        df = df[df["Material"].notna() & df["Documento compras"].notna()].copy()

        if df.empty:
            return pd.DataFrame()

        if "Posición" in df.columns:
            df["Posición"] = df["Posición"].apply(_norm_id)

        for date_col in ("Inicio Validez", "Fecha Término", "Fecha Documento"):
            if date_col in df.columns:
                df[date_col] = pd.to_datetime(df[date_col], errors="coerce")

        df["Operación"] = operation
        df["Fuente"] = "ME3L"
        return df.reset_index(drop=True)
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# OC loader
# ---------------------------------------------------------------------------

def load_oc(inputs_root: Path, operation: str) -> pd.DataFrame:
    """Load and normalize ordenes-compra input files for the given operation."""
    try:
        raw_df, _ = consolidate_purchase_orders(inputs_root, [operation.upper()])
    except Exception:
        return pd.DataFrame()

    if raw_df.empty:
        return pd.DataFrame()

    rename_map = {k: v for k, v in _OC_COL_MAP.items() if k in raw_df.columns}
    df = raw_df.rename(columns=rename_map).copy()

    df["Material"] = df["Material"].apply(_norm_material)
    df = df[df["Material"].notna()].copy()

    df["Operación"] = operation.upper()
    df["Fuente"] = "OC"
    return df.reset_index(drop=True)
