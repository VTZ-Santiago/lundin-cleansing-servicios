"""Generate Excel duplicate-materials report."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.dup_check.detector import REPORT_COLUMNS

_HDR_DARK = "1E3A5F"
_HDR_WHITE = "FFFFFF"
_DATE_COLS = frozenset({"Inicio Validez", "Fecha Término", "Fecha Documento", "Fecha Entrega"})
_ROW_COLORS = ("DBEAFE", "FEF9C3")   # alternating light blue / yellow per material group
_SUMMARY_HI = "FFF9C4"


def _hdr_font() -> Font:
    return Font(bold=True, color=_HDR_WHITE)


def _hdr_fill() -> PatternFill:
    return PatternFill(start_color=_HDR_DARK, end_color=_HDR_DARK, fill_type="solid")


def _row_fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _safe(v):
    """Convert value to Excel-safe Python type."""
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().date()
    if hasattr(v, "item"):
        return v.item()
    return v


def _autofit(ws, max_width: int = 45) -> None:
    for col_cells in ws.columns:
        width = 10
        for cell in col_cells:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width


# ---------------------------------------------------------------------------
# Sheet writers
# ---------------------------------------------------------------------------

def _write_detail_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    """One row per duplicate record, alternating row color per material group."""
    ws = wb.create_sheet(title=title)
    headers = list(df.columns)
    date_indices = {i + 1 for i, h in enumerate(headers) if h in _DATE_COLS}
    mat_idx = headers.index("Material") if "Material" in headers else None

    # Header row
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = _hdr_font()
        cell.fill = _hdr_fill()
        cell.alignment = Alignment(wrap_text=False)

    if df.empty:
        ws.cell(row=2, column=1, value="(Sin duplicados encontrados)")
    else:
        mat_color_map: dict[str, str] = {}
        color_counter = 0
        for row_idx, row_tuple in enumerate(df.itertuples(index=False, name=None), 2):
            material = row_tuple[mat_idx] if mat_idx is not None else None
            if material not in mat_color_map:
                mat_color_map[material] = _ROW_COLORS[color_counter % len(_ROW_COLORS)]
                color_counter += 1
            fill = _row_fill(mat_color_map[material])
            for col_idx, raw_val in enumerate(row_tuple, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                sv = _safe(raw_val)
                cell.value = sv
                cell.fill = fill
                if col_idx in date_indices and sv is not None:
                    cell.number_format = "DD-MM-YYYY"

    ws.freeze_panes = "A2"
    if not df.empty:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df) + 1}"
    ws.row_dimensions[1].height = 28
    _autofit(ws)


def _write_summary_sheet(
    wb: Workbook,
    source_counts: dict[tuple[str, str], int],
    dup_stats: dict[tuple[str, str], dict],
) -> None:
    ws = wb.active
    ws.title = "Resumen"

    title_font = Font(bold=True, size=13, color=_HDR_DARK)
    ws.cell(row=1, column=1,
            value="MATERIALES DUPLICADOS EN MÚLTIPLES CONTRATOS/OC ACTIVOS").font = title_font
    ws.cell(row=2, column=1,
            value=f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    ws.cell(row=3, column=1,
            value="Criterio: mismo material en >1 documento de compra activo "
                  "(Indicador de borrado vacío)")

    headers = [
        "Operación", "Fuente",
        "Registros Activos", "Materiales Duplicados",
        "Documentos Afectados", "Filas en Reporte",
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = _hdr_font()
        cell.fill = _hdr_fill()
        cell.alignment = Alignment(wrap_text=False)

    hi_fill = PatternFill(start_color=_SUMMARY_HI, end_color=_SUMMARY_HI, fill_type="solid")
    row_idx = 6
    for (op, fuente), total_active in sorted(source_counts.items()):
        stats = dup_stats.get((op, fuente), {})
        dup_mats = stats.get("dup_materials", 0)
        values = [
            op, fuente, total_active,
            dup_mats,
            stats.get("dup_docs", 0),
            stats.get("dup_rows", 0),
        ]
        for col_idx, v in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=v)
            if dup_mats > 0:
                cell.fill = hi_fill
        row_idx += 1

    col_widths = [14, 10, 22, 26, 26, 20]
    for col_idx, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.row_dimensions[5].height = 24


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate_report(
    source_counts: dict[tuple[str, str], int],
    dup_dfs: dict[tuple[str, str], pd.DataFrame],
    output_path: Path,
) -> Path:
    """Build Excel report with one summary sheet and one detail sheet per (operation, source).

    source_counts: (operation, source) → count of active records
    dup_dfs:       (operation, source) → DataFrame of duplicate rows
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    dup_stats: dict[tuple[str, str], dict] = {}
    for key, df in dup_dfs.items():
        if df.empty:
            dup_stats[key] = {"dup_materials": 0, "dup_docs": 0, "dup_rows": 0}
        else:
            dup_stats[key] = {
                "dup_materials": int(df["Material"].nunique()),
                "dup_docs": int(df["Documento compras"].nunique())
                if "Documento compras" in df.columns else 0,
                "dup_rows": len(df),
            }

    _write_summary_sheet(wb, source_counts, dup_stats)

    # Detail sheets — fixed order: OC then ME3L, MLCC then CCMC
    _SHEET_ORDER = [
        ("MLCC", "OC",   "OC_MLCC"),
        ("CCMC", "OC",   "OC_CCMC"),
        ("MLCC", "ME3L", "ME3L_MLCC"),
        ("CCMC", "ME3L", "ME3L_CCMC"),
    ]
    for op, src, label in _SHEET_ORDER:
        df = dup_dfs.get((op, src), pd.DataFrame(columns=REPORT_COLUMNS))
        _write_detail_sheet(wb, label, df)

    # Combined "Todos" sheet — filter out frames that are fully empty to avoid FutureWarning
    all_frames = [df for df in dup_dfs.values() if not df.empty and len(df.columns) > 0]
    if all_frames:
        combined = (
            pd.concat([df.dropna(axis=1, how="all") for df in all_frames],
                      ignore_index=True, sort=False)
            .sort_values(["Operación", "Fuente", "Material", "Documento compras"])
            .reset_index(drop=True)
        )
        for col in REPORT_COLUMNS:
            if col not in combined.columns:
                combined[col] = None
        combined = combined[REPORT_COLUMNS]
    else:
        combined = pd.DataFrame(columns=REPORT_COLUMNS)
    _write_detail_sheet(wb, "Todos", combined)

    wb.save(str(output_path))
    return output_path
