import openpyxl
import os
from datetime import datetime

files_to_verify = [
    "outputs/ordenes_compra/consolidado/po_consolidado_MLCC_2025plus.xlsx",
    "outputs/ordenes_compra/consolidado/po_consolidado_CCMC_2025plus_actualizado.xlsx"
]
stale_check = "outputs/ordenes_compra/consolidado/po_consolidado_CCMC_2025plus.xlsx"

def verify(filepath):
    print(f"\n--- File: {filepath} ---")
    if not os.path.exists(filepath):
        print("Exists: No")
        return
    print("Exists: Yes")
    
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        print(f"Sheet names: {sheet_names}")
        
        ws = wb.active # Assuming the first/active sheet
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            print("Empty sheet.")
            return
        
        header = rows[0]
        print(f"Header: {header}")
        
        data = rows[1:]
        data_count = len(data)
        print(f"Data row count: {data_count}")
        
        # Identify column indices
        try:
            h_list = [str(h).strip() if h is not None else "" for h in header]
            idx_fecha = h_list.index("Fecha de Entrega")
            idx_cant = h_list.index("Por entregar (cantidad)")
            idx_valor = h_list.index("Por entregar (valor)")
            idx_pr = h_list.index("PR/SOLPED")
        except ValueError as e:
            print(f"Error finding columns: {e}")
            return

        invalid_fecha = 0
        count_cant = 0
        count_valor = 0
        count_pr = 0
        
        threshold_date = datetime(2025, 1, 1)
        
        for row in data:
            if not row: continue
            
            # Fecha check
            f = row[idx_fecha]
            if f is None:
                invalid_fecha += 1
            else:
                if isinstance(f, str):
                    try:
                        f = datetime.strptime(f, "%Y-%m-%d") # basic format check
                    except:
                        pass # keep as is or handle
                
                if isinstance(f, datetime):
                    if f < threshold_date:
                        invalid_fecha += 1
                else:
                    invalid_fecha += 1 # Not a valid date or wrong format
            
            # Non-empty checks
            if row[idx_cant] is not None and str(row[idx_cant]).strip() != "":
                count_cant += 1
            if row[idx_valor] is not None and str(row[idx_valor]).strip() != "":
                count_valor += 1
            if row[idx_pr] is not None and str(row[idx_pr]).strip() != "":
                count_pr += 1
                
        print(f"Invalid row count (Fecha < 2025 or missing): {invalid_fecha}")
        print(f"Non-empty 'Por entregar (cantidad)': {count_cant}")
        print(f"Non-empty 'Por entregar (valor)': {count_valor}")
        print(f"Non-empty 'PR/SOLPED': {count_pr}")

    except Exception as e:
        print(f"Error processing workbook: {e}")

for f in files_to_verify:
    verify(f)

print(f"\n--- Stale Check ---")
if os.path.exists(stale_check):
    print(f"{stale_check} exists (potentially stale).")
else:
    print(f"{stale_check} does not exist.")
