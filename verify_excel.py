import openpyxl

def verify_workbook(file_path):
    print(f"\n--- Verifying {file_path} ---")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("Empty workbook")
        return
        
    header = rows[0]
    data = rows[1:]
    
    col_map = {name: i for i, name in enumerate(header) if name}
    print(f"Row count (excluding header): {len(data)}")
    
    # Task 3: Fecha formats
    # Note: data_only=True doesn't preserve number_format if we use values_only=True. 
    # Must get cell objects for number_format.
    
    # Reload without values_only for formats or just use ws.cell
    def get_first_non_empty_format(col_name):
        if col_name not in col_map:
            return "Column not found"
        idx = col_map[col_name] + 1 # 1-based for openpyxl
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=idx)
            if cell.value is not None:
                return cell.number_format
        return "No data"

    if 'Fecha documento' in col_map:
        print(f"Fecha documento format: {get_first_non_empty_format('Fecha documento')}")
    if 'Fecha de Entrega' in col_map:
        print(f"Fecha de Entrega format: {get_first_non_empty_format('Fecha de Entrega')}")

    # Specific checks for CCMC
    if "CCMC" in file_path:
        search_val = "4400001549"
        doc_col = col_map.get('Documento compras')
        cl_col = col_map.get('Cl Docto compras')
        
        found_cl = None
        if doc_col is not None and cl_col is not None:
            for row in data:
                if str(row[doc_col]) == search_val:
                    found_cl = row[cl_col]
                    break
        print(f"Documento {search_val} Cl Docto compras: {found_cl}")
        
        if cl_col is not None:
            unique_cl = sorted(list(set(str(row[cl_col]) for row in data if row[cl_col] is not None)))
            print(f"Unique sample values from Cl Docto compras (up to 10): {unique_cl[:10]}")

verify_workbook('outputs/ordenes_compra/consolidado/po_consolidado_CCMC_2025plus.xlsx')
verify_workbook('outputs/ordenes_compra/consolidado/po_consolidado_MLCC_2025plus.xlsx')
